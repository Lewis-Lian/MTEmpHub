from functools import wraps
from io import BytesIO
from datetime import datetime
from pathlib import Path

from flask import Blueprint, g, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import xlwt
from sqlalchemy.exc import IntegrityError

from models import db
from models.meal_ticket import MealTicketBatch, MealTicketImport, MealTicketImportRow
from models.account_set import AccountSet
from routes.auth_helpers import admin_required, page_permission_required
from services.meal_ticket_service import (
    MealError, begin_write, shift_month, generate, batch_for_write, confirm, adjustment,
    payment, serialize_batch, source_snapshot, participation, unconfirm,
    attendance_recalculation_preview, recalculate_attendance, supplement_person,
)
from services.meal_ticket_import_service import preview, serialize_import, confirm_import, import_rows
from services.meal_ticket_reconciliation import reconcile

meal_tickets_bp = Blueprint('meal_tickets', __name__, url_prefix='/api/meal-tickets')


def handled(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            if request.is_json and not isinstance(request.get_json(silent=True), dict):
                raise MealError('请求内容必须是对象')
            return fn(*args, **kwargs)
        except MealError as exc:
            db.session.rollback()
            return jsonify({'error':str(exc)}), exc.status
        except IntegrityError:
            db.session.rollback()
            return jsonify({'error':'数据冲突，请刷新后重试'}), 409
    return wrapper


def accessible():
    if g.current_user.role == 'admin':
        return None
    from routes.query_core import _accessible_emp_ids
    return _accessible_emp_ids()


def selected_batch():
    recharge_month = request.args.get('recharge_month', '')
    month = shift_month(recharge_month, -1)
    return MealTicketBatch.query.filter_by(month=month).first()


@meal_tickets_bp.get('')
@page_permission_required('meal_ticket_query')
@handled
def get_batch():
    from services.meal_ticket_progress_service import start_progress
    token = request.headers.get('X-Meal-Progress-Token')
    progress = start_progress(token, g.current_user.id) if token else None
    try:
        batch = selected_batch()
        result = serialize_batch(batch, accessible(), check_source=g.current_user.role == 'admin', progress=progress) if batch else None
        response = jsonify(result)
        if progress:
            progress('月度核算数据已加载', 1, 1, status='completed')
        return response
    except Exception:
        if progress:
            progress('月度核算数据读取失败', status='failed')
        raise


@meal_tickets_bp.get('/progress')
@page_permission_required('meal_ticket_query')
@handled
def get_read_progress():
    from services.meal_ticket_progress_service import read_progress
    response = jsonify(read_progress(request.args.get('token'), g.current_user.id))
    response.headers['Cache-Control'] = 'no-store'
    return response


@meal_tickets_bp.post('/generate')
@admin_required
@handled
def generate_batch():
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = generate(body.get('recharge_month'), operator)
    result = serialize_batch(batch)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/confirm')
@admin_required
@handled
def confirm_batch():
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    confirm(batch, operator)
    result = serialize_batch(batch)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/unconfirm')
@admin_required
@handled
def unconfirm_batch():
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    unconfirm(batch)
    result = serialize_batch(batch, check_source=True)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/adjustments')
@admin_required
@handled
def add_adjustment():
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    adjustment(batch, body.get('item_id'), body.get('amount'), body.get('reason'), operator)
    result = serialize_batch(batch)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/participation')
@admin_required
@handled
def set_participation():
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    participation(batch, body.get('item_id'), body.get('excluded'), body.get('reason'), operator)
    result = serialize_batch(batch)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/attendance-recalculation/preview')
@admin_required
@handled
def preview_attendance_recalculation():
    body = request.get_json(silent=True) or {}
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    return jsonify(attendance_recalculation_preview(batch))


@meal_tickets_bp.post('/attendance-recalculation')
@admin_required
@handled
def apply_attendance_recalculation():
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    recalculate_attendance(batch, body.get('source_digest'), g.current_user.username)
    result = serialize_batch(batch, check_source=True)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/supplement-person')
@admin_required
@handled
def supplement_batch_person():
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = batch_for_write(body.get('batch_id'), body.get('version'))
    supplement_person(batch, body.get('emp_id'), body.get('source_digest'),
                      body.get('reason'), g.current_user.username)
    result = serialize_batch(batch, check_source=True)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/payments')
@admin_required
@handled
def add_payment():
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    batch = payment(body, operator)
    result = serialize_batch(batch)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.post('/reconcile')
@admin_required
@handled
def reconcile_batch():
    batch, report = reconcile(request.get_json(silent=True) or {}, g.current_user.username)
    result = serialize_batch(batch)
    result['reconciliation'] = report
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.get('/export-recharge')
@page_permission_required('meal_ticket_query')
@handled
def export_recharge():
    batch = selected_batch()
    if not batch:
        raise MealError('没有核算数据', 404)
    if batch.status != 'confirmed':
        raise MealError('请先确认核算再导出充值表', 409)
    data = serialize_batch(batch, accessible())
    book = xlwt.Workbook()
    sheet = book.add_sheet('充值表')
    row = 0
    for item in data['items']:
        amount = round(item['difference'], 2)
        if amount <= 0:
            continue
        sheet.write(row, 0, item['emp_no'])
        sheet.write(row, 1, amount)
        row += 1
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f'菜票充值_{batch.recharge_month}.xls',
        mimetype='application/vnd.ms-excel')


@meal_tickets_bp.get('/export')
@page_permission_required('meal_ticket_query')
@handled
def export_batch():
    batch = selected_batch()
    if not batch:
        raise MealError('没有核算数据', 404)
    data = serialize_batch(batch, accessible())
    book = Workbook()
    for index, (name, headers, rows) in enumerate((
        ('人员明细', ['工号','姓名','核算部门','实际打卡天数','基础金额','额外补扣','应发金额','净已发金额','差额','核对说明'],
         [[i['emp_no'],i['name'],i['dept_name'],i['days'],i['base_amount'],i['adjustment_amount'],i['due_amount'],i['paid_amount'],i['difference'],
           '；'.join(filter(None, [i['error'], ('本月不发：' if i['excluded'] else '恢复核算：') + i['participation_history'][-1]['reason'] if i['participation_history'] else '']))] for i in data['items']]),
        ('部门汇总', ['部门','人数','基础金额','额外补扣','应发金额','净已发金额','差额'],
         [[d[k] for k in ('dept_name','count','base_amount','adjustment_amount','due_amount','paid_amount','difference')] for d in data['departments']]),
    )):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = name
        sheet.append([f"考勤 {batch.month} / 计划充值 {batch.recharge_month} / {'草稿' if batch.status == 'draft' else '已确认'}"])
        sheet.append(['生成时间', datetime.now().isoformat(timespec='seconds')])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        # Prevent user-entered text becoming spreadsheet formulas.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(('=','+','-','@')):
                    cell.data_type = 's'
        for cell in sheet[3]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245B78')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 20
        sheet.freeze_panes = 'D4' if index == 0 else 'B4'
        sheet.auto_filter.ref = f'A3:{sheet.cell(sheet.max_row,len(headers)).coordinate}'
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f'菜票_{batch.recharge_month}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@meal_tickets_bp.get('/imports')
@admin_required
@handled
def list_imports():
    records = MealTicketImport.query.order_by(MealTicketImport.id.desc()).all()
    return jsonify([serialize_import(r) for r in records])


@meal_tickets_bp.post('/imports')
@admin_required
@handled
def preview_import():
    operator = g.current_user.username
    file = request.files.get('file')
    if not file:
        raise MealError('请选择 xlsx 文件')
    payload = file.read(20 * 1024 * 1024 + 1)
    begin_write()
    record = preview(payload, file.filename or '', request.form.get('month'), request.form.get('month_kind'), operator)
    result = serialize_import(record)
    db.session.commit()
    return jsonify(result)


@meal_tickets_bp.get('/imports/<int:identifier>/comparison')
@admin_required
@handled
def import_comparison(identifier):
    record = db.session.get(MealTicketImport, identifier)
    if not record:
        raise MealError('历史原账不存在', 404)
    has_account = AccountSet.query.filter_by(month=record.month).first() is not None
    sources = {row['emp_id']:row for row in source_snapshot(record.month)} if has_account else {}
    result = []
    for row in import_rows(record):
        if row.data['kind'] != 'person' or row.data.get('skip'):
            continue
        source = sources.get(row.emp_id)
        error = source['error'] if source else '缺少对应考勤账套或人员依据'
        amount = source['base_cents'] / 100 if source and not error else None
        old = row.data['amount']
        result.append({'id':row.id, 'emp_no':row.data['emp_no'], 'name':row.data['name'],
            'historical_amount':old, 'days':source['days'] if source else None,
            'base_amount':amount, 'difference':round(old - amount, 2) if old is not None and amount is not None else None,
            'error':error})
    return jsonify(result)


@meal_tickets_bp.delete('/imports/<int:identifier>')
@admin_required
@handled
def import_cancel(identifier):
    begin_write()
    record = MealTicketImport.query.filter_by(id=identifier).with_for_update().first()
    if not record:
        raise MealError('导入预览不存在', 404)
    if record.status != 'preview':
        raise MealError('已确认入账的历史记录不能取消', 409)
    path = Path(record.stored_path)
    MealTicketImportRow.query.filter_by(import_key=record.key).delete(synchronize_session=False)
    db.session.delete(record)
    db.session.commit()
    path.unlink(missing_ok=True)
    return jsonify({'deleted': identifier})


@meal_tickets_bp.post('/imports/<int:identifier>/confirm')
@admin_required
@handled
def import_confirm(identifier):
    operator = g.current_user.username
    body = request.get_json(silent=True) or {}
    begin_write()
    record = MealTicketImport.query.filter_by(id=identifier).with_for_update().first()
    if not record:
        raise MealError('导入预览不存在', 404)
    record = confirm_import(record, body.get('rows'), operator)
    result = serialize_import(record)
    db.session.commit()
    return jsonify(result)
