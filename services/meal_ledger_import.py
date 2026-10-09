"""Server-owned previews for historical business ledgers."""
from calendar import monthrange
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import hashlib
import re
import uuid
import zipfile

from flask import current_app
from openpyxl import load_workbook

from models import db
from models.meal_ledger import MealLedgerImport, MealLedgerRecord
from services import meal_ledger_service as ledger
from services.meal_ticket_service import MealError, required_text

MAX_SIZE = 10 * 1024 * 1024


def identifier(value):
    if value is None:
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def serialized(record):
    return {'id': record.id, 'kind': record.kind, 'month': record.month,
        'filename': record.source_filename, 'status': record.status, 'rows': record.data['rows']}


def suspected_duplicate(kind, body):
    if kind not in ('external', 'clearance'):
        return False
    month, _, amount, data, _ = ledger.validated(kind, body)
    for existing in ledger.records(kind, month):
        if existing.voided or existing.amount_cents != amount:
            continue
        if data['emp_no'] and existing.data.get('emp_no') and data['emp_no'] != existing.data['emp_no']:
            continue
        if kind == 'external' and data['category'] != existing.data.get('category'):
            continue
        if existing.data.get('name') == data['name']:
            return True
    return False


def preview(kind, month, upload, operator):
    if kind not in ledger.KINDS:
        raise MealError('导入类型无效')
    ledger.month_checked(month)
    raw = upload.read(MAX_SIZE + 1)
    if len(raw) > MAX_SIZE:
        raise MealError('文件不能超过 10 MB')
    checksum = hashlib.sha256(raw).hexdigest()
    old = MealLedgerImport.query.filter_by(kind=kind, month=month, file_digest=checksum).first()
    if old:
        return old
    try:
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise MealError('解压后的表格过大')
        formulas = load_workbook(BytesIO(raw), data_only=False, read_only=True)
        values = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except (ValueError, OSError, zipfile.BadZipFile, KeyError):
        raise MealError('请上传有效的 xlsx 文件') from None
    rows = []
    for sheet in formulas:
        if sheet.max_row > 20000 or sheet.max_column > 100:
            raise MealError('工作表范围过大')
        header = None
        headers = []
        blocks = [0]
        title_years = set()
        for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            for value in row:
                title_years.update(re.findall(r'(20\d{2})年', str(value or '')))
            normalized = [re.sub(r'\s+', '', str(v or '')) for v in row]
            if kind == 'department' and '部门' in normalized and '金额' in normalized:
                header, headers = index, normalized
                blocks = [i for i, value in enumerate(headers) if value == '部门'][:2]
                break
            if ((kind == 'clearance' and '人员姓名' in normalized[:8] and any(label in normalized[:8] for label in ('卡余额', '取款金额(元)', '取款金额（元）')))
                    or (kind == 'external' and '姓名' in normalized[:12] and '发放金额（元）' in normalized[:12] and '类别' in normalized[:12])):
                header, headers = index, normalized
                if kind == 'external' and '卡号' in headers:
                    last = headers.index('卡号') + 1
                    headers = headers[:last + 1]
                    if len(headers) > last and not headers[last]:
                        headers[last] = '备注'
                # The reference workbook puts several pivots beside the source table.
                seen_labels = set()
                for column, label in enumerate(headers):
                    if label and label in seen_labels:
                        headers = headers[:column]
                        break
                    if label:
                        seen_labels.add(label)
                break
            if kind == 'consumption' and '二楼消费金额' in normalized and '三楼消费金额' in normalized:
                header, headers = index, normalized
                break
            if index >= 8:
                break
        if header is None:
            continue
        cache = values[sheet.title]
        cached_rows = cache.iter_rows(min_row=header + 1, values_only=True)
        for row_number, (formula, cached) in enumerate(zip(sheet.iter_rows(min_row=header + 1, values_only=True), cached_rows), header + 1):
            for block in blocks:
                h = headers[block:block + 4] if kind == 'department' else headers
                f = formula[block:block + 4] if kind == 'department' else formula[:len(headers)]
                c = cached[block:block + 4] if kind == 'department' else cached[:len(headers)]
                source = dict(zip(h, c))
                def get(*labels):
                    return next((source[label] for label in labels if label in source), None)
                name = get('人员姓名', '姓名')
                dept = get('部门名称', '部门', '部 门')
                amount = get('取款金额(元)', '取款金额（元）', '卡余额', '发放金额（元）', '金额')
                if kind == 'department' and (not dept or '合计' in str(dept) or '总计' in str(dept)):
                    continue
                if kind in ('clearance', 'external') and (not name and not amount):
                    continue
                if kind in ('clearance', 'external') and any(re.fullmatch(r'(总合计|合计|总计)[：:\s]*', str(v or '').strip()) for v in (name, dept, get('外来人员（部门）'))):
                    continue
                if kind == 'consumption' and not re.fullmatch(r'\d{1,2}月', str(get('月份') or '')):
                    continue
                sheet_month = re.search(r'(\d{1,2})月', sheet.title)
                value_month = re.search(r'(\d{1,2})月', str(get('月份') or ''))
                month_number = int((value_month or sheet_month).group(1)) if value_month or sheet_month else int(month[5:])
                sheet_year = re.search(r'(20\d{2})年', sheet.title)
                proposed = f'{sheet_year.group(1) if sheet_year else month[:4]}-{month_number:02d}'
                item = {'index': len(rows), 'sheet': sheet.title, 'row': row_number, 'block': block,
                    'month': proposed, 'date': '', 'emp_no': identifier(get('人员编号')),
                    'name': identifier(name), 'dept_name': identifier(dept),
                    'card_no': identifier(get('卡号')), 'remark': identifier(get('备注')),
                    'registrar': identifier(get('填表人', '登记人')), 'amount': amount,
                    'unit': identifier(get('外来人员（部门）')), 'period': identifier(get('时间')),
                    'days': identifier(get('天数')), 'category': 'card' if get('类别') == '充卡' else 'paper' if get('类别') == '纸质' else '',
                    'floor2': get('二楼消费金额'), 'floor3': get('三楼消费金额'),
                    'historical_amount': amount if kind == 'department' else None,
                    'historical_recharge': get('充值金额'), 'historical_consumption': get('消费金额'), 'historical_recovered': get('收回金额'),
                    'raw': [identifier(v) for v in f], 'error': '', 'warning': ''}
                recorded_date = get('日期', '处理日期')
                try:
                    if isinstance(recorded_date, (date, datetime)):
                        item['date'] = recorded_date.strftime('%Y-%m-%d')
                    elif isinstance(recorded_date, (int, float)):
                        from openpyxl.utils.datetime import from_excel
                        parsed = from_excel(recorded_date, formulas.epoch)
                        if not isinstance(parsed, (date, datetime)):
                            raise MealError('原表日期无效，请修正处理日期')
                        item['date'] = parsed.strftime('%Y-%m-%d')
                    elif recorded_date:
                        item['date'] = date.fromisoformat(str(recorded_date).replace('/', '-')).isoformat()
                    ledger.month_checked(proposed)
                    default_date = date(int(proposed[:4]), month_number, monthrange(int(proposed[:4]), month_number)[1])
                    if not item['date']:
                        item['date'] = default_date.isoformat()
                    ledger.validated(kind, item)
                except MealError as exc:
                    item['error'] = str(exc)
                except (ValueError, TypeError, OverflowError):
                    item['error'] = '原表日期无效，请修正处理日期'
                if any(isinstance(v, str) and v.startswith('=') and cached_value is None for v, cached_value in zip(f, c)):
                    item['warning'] = '存在无缓存公式，请核对金额；缺失金额必须修正后导入'
                if title_years and title_years != {proposed[:4]}:
                    item['warning'] += '；标题年份与所选年份不一致，请确认实际年月'
                if not recorded_date and kind in ('external', 'clearance'):
                    item['warning'] += '；原表无处理日期，建议月末日期，请核实后确认'
                if not item['error'] and suspected_duplicate(kind, item):
                    item['warning'] += '；与已有记录疑似重复，请确认是否为另一笔真实记录'
                    item['duplicate'] = True
                rows.append(item)
    formulas.close()
    values.close()
    if not rows:
        raise MealError('未找到所选台账类型的有效表头和明细')
    root = Path(current_app.config['UPLOAD_FOLDER']) / 'meal_ledger_sources'
    root.mkdir(parents=True, exist_ok=True)
    path = root / (uuid.uuid4().hex + '.xlsx')
    path.write_bytes(raw)
    record = MealLedgerImport(kind=kind, month=month, source_filename=Path(upload.filename or '导入.xlsx').name[:255],
        file_digest=checksum, stored_path=str(path), data={'rows': rows}, operator=operator)
    db.session.add(record)
    db.session.flush()
    return record


def confirm(record, choices, operator):
    if record.status == 'confirmed':
        return record
    if not isinstance(choices, list) or len(choices) != len(record.data['rows']):
        raise MealError('请逐行确认月份、处理日期，或明确跳过')
    selected = {}
    for choice in choices:
        if not isinstance(choice, dict) or type(choice.get('index')) is not int or choice['index'] in selected:
            raise MealError('导入行选择无效')
        selected[choice['index']] = choice
    saved = []
    for item in record.data['rows']:
        choice = selected.get(item['index'])
        if not choice:
            raise MealError('缺少导入行确认')
        if choice.get('skip') is True:
            saved.append({**item, 'skipped': True})
            continue
        month = ledger.month_checked(choice.get('month'))
        day = required_text(choice.get('date'), '处理日期', 10)
        body = {**item, **{k: choice[k] for k in ('amount', 'floor2', 'floor3', 'category', 'dept_name', 'name', 'historical_amount') if k in choice},
            'month': month, 'date': day, 'request_key': f'ledger-import:{record.key}:{item["index"]}'}
        if suspected_duplicate(record.kind, body) and choice.get('accept_duplicate') is not True:
            raise MealError(f'第 {item["index"] + 1} 行与已登记记录疑似重复，请确认另一笔或跳过')
        source = f'xlsx:{record.kind}:{record.file_digest}:{item["sheet"]}:{item["row"]}:{item["block"]}'
        # Sheet names can be long; the identity remains bounded and stable.
        source = 'xlsx:' + ledger.digest(source)
        created = ledger.create(record.kind, body, operator, source)
        created.data = {**created.data, 'import_key': record.key, 'source_sheet': item['sheet'], 'source_row': item['row']}
        saved.append({**body, 'error': '', 'skipped': False})
    record.data = {'rows': saved}
    record.status = 'confirmed'
    db.session.flush()
    return record
