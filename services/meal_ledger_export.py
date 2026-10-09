"""Archive workbooks contain posted values rather than legacy calculation formulas."""
from datetime import datetime
from io import BytesIO
from math import ceil
from unicodedata import east_asian_width

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models.meal_ticket import MealTicketBatch
from services import meal_ledger_service as ledger
from services.meal_ticket_service import MealError, serialize_batch


def sheet(book, name, title, headers, rows, status='有效记录'):
    ws = book.create_sheet(name)
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.append([f'{status} · 导出时间 {datetime.now().isoformat(timespec="seconds")}'])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[3].height = 28
    ws['A1'].font = Font(size=15, bold=True)
    border = Border(bottom=Side(style='hair', color='D9E2EA'))
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = border
            if cell.row == 3:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='245B78')
            elif isinstance(cell.value, (int, float)):
                cell.number_format = '0.00'
            elif isinstance(cell.value, str) and cell.value.startswith(('=', '+', '-', '@')):
                cell.data_type = 's'
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = 24 if header in ('备注', '部门名称', '外来人员（部门）') else 18
    for row in ws.iter_rows(min_row=4):
        lines = 1
        for cell in row:
            if isinstance(cell.value, str):
                width = ws.column_dimensions[cell.column_letter].width - 2
                count = sum(max(1, ceil(sum(2 if east_asian_width(c) in ('W', 'F') else 1 for c in part) / width))
                            for part in cell.value.split('\n'))
                lines = max(lines, count)
        ws.row_dimensions[row[0].row].height = 16 * lines
    ws.freeze_panes = 'D4'
    ws.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{ws.max_row}'
    ws.print_title_rows = '1:3'
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.print_area = f'A1:{get_column_letter(len(headers))}{ws.max_row}'
    return ws


def export(report, month=None, year=None, accessible=None, filters=None):
    book = Workbook()
    book.remove(book.active)
    filters = filters or {}
    def matches(item):
        return (not filters.get('dept_name') or item['dept_name'] == filters['dept_name']) and (
            not filters.get('search') or filters['search'].casefold() in (item['emp_no'] + item['name']).casefold()) and (
            filters.get('person_type') not in ('employee', 'manager') or item['is_manager'] == (filters['person_type'] == 'manager'))
    if report == 'annual':
        data = ledger.annual(year)
        headers = ['月份', '充值金额', '消费金额', '二楼消费金额', '三楼消费金额', '收回金额', '纸质领用金额', '消费数据状态', '历史充值原额', '历史收回原额']
        rows = [[r['month'], r['recharge_amount'], r['consumption_amount'], r['floor2_amount'], r['floor3_amount'], r['recovered_amount'], r['paper_amount'],
            '已录入' if r['consumption_amount'] is not None else '未录入', r['historical_recharge_amount'], r['historical_recovered_amount']] for r in data['months']]
        total = data['totals']
        rows.append(['总合计', total['recharge_amount'], total['consumption_amount'], total['floor2_amount'], total['floor3_amount'], total['recovered_amount'], total['paper_amount'], f'消费已录入 {data["consumption_months"]}/12 月'])
        sheet(book, '全年菜票录入汇总', f'{data["year"]}年各月菜票充值使用汇总', headers, rows)
        period = str(data['year'])
    else:
        yearly = not month and report in ('department', 'external', 'clearance')
        period = str(ledger.annual(year)['year']) if yearly else ledger.month_checked(month)
        periods = [f'{int(period):04d}-{n:02d}' for n in range(1, 13)] if yearly else [month]
        if report == 'recharge':
            batch = MealTicketBatch.query.filter_by(recharge_month=month).first()
            if not batch:
                raise MealError('本月没有核算数据', 404)
            data = serialize_batch(batch, accessible)
            items = [i for i in data['items'] if matches(i)]
            status = '草稿' if batch.status == 'draft' else '已确认核算；实际充值列采用已登记净发放金额'
            title = f'{month}充值记录（考勤 {batch.month}）'
            from routes.query_core import _build_abnormal_rows
            ordinary_ids = [i['emp_id'] for i in items if not i['is_manager']]
            abnormal = _build_abnormal_rows(batch.month, ordinary_ids) if ordinary_ids else []
            counts = {str(r['emp_no']): r.get('abnormal_count', 0) for r in abnormal}
            sheet(book, '员工异常查询1', title, ['部门名称', '人员编号', '人员姓名', '异常考勤次数'],
                [[i['dept_name'], i['emp_no'], i['name'], counts.get(i['emp_no'], 0)] for i in items if not i['is_manager']], status)
            sheet(book, '员工充值记录', title,
                ['部门名称', '人员编号', '人员名称', '考勤天数', '充值天数', '充值金额', '异常刷卡次数', '异常刷卡金额', '额外情况', '实际充值金额', '备注'],
                [[i['dept_name'], i['emp_no'], i['name'], i['days'], i['days'], i['base_amount'], counts.get(i['emp_no'], 0), None,
                    i['adjustment_amount'], i['paid_amount'], '；'.join(filter(None, [i['error'], '异常次数仅展示，不自动扣款', *[a['reason'] for a in i['adjustments']]]))]
                    for i in items if not i['is_manager']], status)
            sheet(book, '管理人员查询', title,
                ['部门', '人员编号', '姓名', '考勤天数', '充值天数', '额外情况', '充值金额', '备注'],
                [[i['dept_name'], i['emp_no'], i['name'], i['days'], i['days'], i['adjustment_amount'], i['paid_amount'],
                    '实际发放金额；' + '；'.join(a['reason'] for a in i['adjustments'])] for i in items if i['is_manager']], status)
        elif report == 'department':
            for selected in periods:
                data = ledger.departments(selected, accessible)
                items = [r for r in data['items'] if not filters.get('dept_name') or r['dept_name'] == filters['dept_name']]
                entries = [[r['dept_name'], r['total_paid_amount'], r['registrar'], r['remark']] for r in items]
                split = (len(entries) + 1) // 2
                rows = [entries[index] + [''] + (entries[index + split] if index + split < len(entries) else ['', '', '', ''])
                        for index in range(split)]
                rows.append(['合计', round(sum(r['total_paid_amount'] for r in items), 2), '', '员工净发放＋客人充卡＋纸质菜票；不含月底清零'])
                ws = sheet(book, selected[5:] + '月', f'{selected}月份部门菜票发放汇总', ['部门', '金额', '登记人', '备注', '', '部门', '金额', '登记人', '备注'], rows,
                    '草稿' if data['status'] == 'draft' else '无核算数据' if data['status'] == 'no_batch' else '已确认核算')
                ws.column_dimensions['E'].width = 3
                ws.auto_filter.ref = None
                ws.freeze_panes = 'A4'
                details = [[r['dept_name'], r['paid_amount'], r['external_card_amount'],
                    r['external_paper_amount'], r['total_paid_amount']] for r in items]
                details.append(['合计', *[round(sum(r[key] for r in items), 2) for key in
                    ('paid_amount', 'external_card_amount', 'external_paper_amount', 'total_paid_amount')]])
                sheet(book, selected[5:] + '月明细', f'{selected}月份部门菜票发放明细',
                    ['部门', '员工净实际发放', '客人卡充值', '客人纸质菜票', '部门发放合计'], details)
        elif report in ('external', 'clearance'):
            rows = [r for r in ledger.records(report, month=None if yearly else month, year=period if yearly else None) if not r.voided]
            if report == 'external':
                sheet(book, '外来人员菜票领用', f'{period}外来人员纸质和充值菜票领用记录',
                    ['月份', '日期', '填表人', '部门', '外来人员（部门）', '姓名', '时间', '天数', '发放金额（元）', '类别', '卡号', '备注'],
                    [[r.month, r.record_date.isoformat(), r.data.get('registrar') or r.operator, r.data['dept_name'], r.data['unit'], r.data['name'], r.data['period'],
                        r.data['days'], r.amount_cents / 100, '充卡' if r.data['category'] == 'card' else '纸质', r.data['card_no'], r.data['remark']] for r in rows])
                groups = {}
                for r in rows:
                    key = (r.month, r.data['dept_name'], r.data['category'])
                    groups[key] = groups.get(key, 0) + r.amount_cents
                sheet(book, '领用分类汇总', f'{period}领用分类汇总', ['月份', '部门', '类别', '发放金额（元）'],
                    [[selected, dept, '充卡' if category == 'card' else '纸质', value / 100] for (selected, dept, category), value in sorted(groups.items())])
            else:
                for selected in periods:
                    sheet(book, selected[5:] + '月', f'{selected}月末清零取款明细', ['人员编号', '人员姓名', '卡号', '部门名称', '取款金额(元)', '处理日期', '备注'],
                        [[r.data['emp_no'], r.data['name'], r.data['card_no'], r.data['dept_name'], r.amount_cents / 100, r.record_date.isoformat(), r.data['remark']] for r in rows if r.month == selected])
        else:
            raise MealError('报表类型无效')
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output, f'菜票_{report}_{period}.xlsx'
