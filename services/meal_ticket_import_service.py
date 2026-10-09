"""Read historical workbooks into a separate, reviewable ledger."""
from io import BytesIO
from pathlib import Path
import hashlib
import re
import uuid
import zipfile

from flask import current_app
from openpyxl import load_workbook

from models import db
from models.employee import Employee
from models.system_setting import SystemSetting
from models.meal_ticket import MealTicketImport, MealTicketImportRow
from services.meal_ticket_service import MealError, cents, shift_month, required_text
from sqlalchemy import text


def lock_month(month):
    # All imports for a month share this lock, including when no account-set
    # exists. MySQL's upsert locks the unique key; SQLite uses BEGIN IMMEDIATE.
    key = 'meal_import_lock:' + month
    if db.engine.dialect.name == 'mysql':
        db.session.execute(text("INSERT INTO system_settings (`key`, value) VALUES (:key, '') ON DUPLICATE KEY UPDATE id=id"), {'key':key})
    else:
        if SystemSetting.query.filter_by(key=key).first() is None:
            db.session.add(SystemSetting(key=key, value=''))
            db.session.flush()
    SystemSetting.query.filter_by(key=key).with_for_update().one()


def label(value):
    return re.sub(r'\s+', '', str(value or ''))


def identifier(value):
    return str(int(value)) if isinstance(value, (int, float)) and value == int(value) else str(value or '').strip()


def import_rows(record):
    return MealTicketImportRow.query.filter_by(import_key=record.key).order_by(MealTicketImportRow.id).all()


def serialize_import(record):
    rows = [{'id':row.id, **row.data, 'emp_id':row.emp_id, 'original_emp_id':row.emp_id} for row in import_rows(record)]
    person_total = sum(row['amount'] or 0 for row in rows if row['kind'] == 'person' and not row.get('skip'))
    department_total = sum(row['amount'] or 0 for row in rows if row['kind'] == 'department' and not row.get('skip'))
    grouped = {}
    for row in rows:
        if row.get('skip'):
            continue
        group = grouped.setdefault(row['dept_name'], {'dept_name':row['dept_name'], 'person_amount':0, 'historical_amount':0})
        key = 'person_amount' if row['kind'] == 'person' else 'historical_amount'
        group[key] = round(group[key] + (row['amount'] or 0), 2)
    for row in grouped.values():
        row['difference'] = round(row['person_amount'] - row['historical_amount'], 2)
    return {'id':record.id, 'filename':record.source_filename, 'month':record.month,
            'recharge_month':record.recharge_month, 'status':record.status,
            'rows':rows, 'person_total':round(person_total, 2), 'department_total':round(department_total, 2),
            'departments':list(grouped.values())}


def preview(payload, filename, month, month_kind, operator):
    if month_kind not in ('attendance','recharge'):
        raise MealError('请选择文件月份表示考勤月还是充值月')
    shift_month(month, 0)
    attendance = month if month_kind == 'attendance' else shift_month(month, -1)
    recharge = shift_month(attendance, 1)
    lock_month(attendance)
    if not filename.lower().endswith('.xlsx') or len(payload) > 20 * 1024 * 1024:
        raise MealError('仅支持不超过 20 MiB 的 xlsx')
    checksum = hashlib.sha256(payload).hexdigest()
    old = MealTicketImport.query.filter_by(file_digest=checksum, month=attendance).first()
    if old:
        return old
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise MealError('工作簿解压内容过大')
        formulas = load_workbook(BytesIO(payload), read_only=True, data_only=False)
        values = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except (ValueError, OSError, zipfile.BadZipFile, KeyError) as exc:
        raise MealError('工作簿无法读取') from exc
    people = {e.emp_no:e for e in Employee.query.all()}
    parsed = []
    try:
        for sheet in formulas:
            # Annual department books have one page per month. Only the
            # explicitly selected page belongs to this preview.
            page_month = re.fullmatch(r'(\d{1,2})月', sheet.title)
            if page_month and int(page_month.group(1)) != int(month[5:]):
                continue
            if not ('充值' in sheet.title or '管理人员查询' == sheet.title or page_month or '菜票登记' == sheet.title):
                continue
            header_row, groups = None, []
            metadata = [filename, sheet.title]
            for n, raw in enumerate(sheet.iter_rows(max_row=8, max_col=32, values_only=True), 1):
                metadata.extend(str(x) for x in raw if isinstance(x,str))
                headers = [label(x) for x in raw]
                if '人员编号' in headers:
                    amount = next((x for x in ('实际充值金额','充值金额') if x in headers), None)
                    if amount:
                        groups = [('person', headers.index('人员编号'),
                            next((headers.index(x) for x in ('人员名称','人员姓名','姓名') if x in headers), None),
                            next((headers.index(x) for x in ('部门名称','部门') if x in headers), None), headers.index(amount))]
                else:
                    groups = [('department', None, None, i, i+1) for i,h in enumerate(headers[:-1]) if h == '部门' and headers[i+1] == '金额']
                if groups:
                    header_row = n
                    break
            if header_row is None:
                continue
            years = set(re.findall(r'(?<!\d)(20\d{2})(?!\d)', ' '.join(metadata)))
            months = set(re.findall(r'(?<!\d)(1[0-2]|0?[1-9])\s*月', ' '.join(metadata)))
            conflict = ''
            if any(year != month[:4] for year in years) or any(int(number) != int(month[5:]) for number in months):
                conflict = '文件或页内年月与所选月份不一致，请核对归属月份'
            if sheet.max_row > 10000:
                raise MealError('充值数据页超过 10000 行')
            cached = values[sheet.title].iter_rows(min_row=header_row+1, max_row=sheet.max_row, max_col=32, values_only=True)
            source = sheet.iter_rows(min_row=header_row+1, max_row=sheet.max_row, max_col=32, values_only=True)
            for n, (raw, computed) in enumerate(zip(source, cached), header_row+1):
                for kind, emp_col, name_col, dept_col, amount_col in groups:
                    emp_no = identifier(raw[emp_col]) if emp_col is not None else ''
                    dept_name = (identifier(raw[dept_col]) if dept_col is not None else '') or '未分配部门'
                    person_name = str(raw[name_col] or '').strip() if name_col is not None else ''
                    if kind == 'person' and ((not emp_no and not person_name) or any(x in emp_no + person_name for x in ('合计','总计','汇总'))):
                        continue
                    if kind == 'department' and (not dept_name or any(x in dept_name for x in ('合计','总计','汇总')) or raw[amount_col] is None):
                        continue
                    raw_amount = raw[amount_col]
                    value = computed[amount_col] if isinstance(raw_amount, str) and raw_amount.startswith('=') else raw_amount
                    amount, error = None, ''
                    try:
                        amount = cents(value) / 100
                    except MealError:
                        error = '金额缺失、公式缺少缓存或错误，请填写金额并说明或跳过'
                    emp = people.get(emp_no)
                    if kind == 'person' and not emp:
                        error = error or '工号未匹配，请选择人员或跳过'
                    if len(dept_name) > 100:
                        error = error or '历史部门名称不能超过 100 字'
                    name = person_name
                    if not dept_name:
                        dept_name = '未分配部门'
                    parsed.append({'sheet':sheet.title, 'row':n, 'kind':kind, 'emp_no':emp_no,
                        'name':name, 'dept_name':dept_name, 'amount':amount, 'formula':raw_amount if isinstance(raw_amount,str) and raw_amount.startswith('=') else None,
                        'original_amount':amount, 'original_dept_name':dept_name, 'error':error, 'emp_id':emp.id if emp else None,
                        'skip':False, 'correction_reason':'', 'period_conflict':conflict, 'period_confirmed':False})
    finally:
        formulas.close()
        values.close()
    if not parsed:
        raise MealError('没有可导入的人员充值或部门登记数据；空表、取款和外来人员页不导入')
    seen = set()
    for row in parsed:
        if row['kind'] == 'person':
            if row['emp_no'] in seen:
                row['error'] = '同一月份重复人员，请核对并跳过重复行'
            seen.add(row['emp_no'])
    root = Path(current_app.config['UPLOAD_FOLDER']) / 'meal_tickets'
    root.mkdir(parents=True, exist_ok=True)
    path = root / (uuid.uuid4().hex + '.xlsx')
    path.write_bytes(payload)
    record = MealTicketImport(month=attendance, recharge_month=recharge, file_digest=checksum,
        source_filename=Path(filename).name[:255], stored_path=str(path), operator=operator)
    db.session.add(record)
    db.session.flush()
    for row in parsed:
        emp_id = row.pop('emp_id')
        db.session.add(MealTicketImportRow(import_key=record.key, month=attendance, emp_id=emp_id, data=row))
    db.session.flush()
    return record


def confirm_import(record, submitted, operator):
    lock_month(record.month)
    if record.status == 'confirmed':
        return record
    if not isinstance(submitted, list):
        raise MealError('导入行无效')
    existing = import_rows(record)
    edits = {r.get('id'):r for r in submitted if isinstance(r,dict)}
    if len(edits) != len(existing) or set(edits) != {r.id for r in existing}:
        raise MealError('请完整提交预览行')
    previous = MealTicketImportRow.query.join(MealTicketImport, MealTicketImportRow.import_key == MealTicketImport.key).filter(
        MealTicketImport.month == record.month, MealTicketImport.status == 'confirmed').all()
    seen = {r.emp_id for r in previous if r.data['kind'] == 'person' and not r.data.get('skip')}
    for row in existing:
        edit = edits[row.id]
        if type(edit.get('skip')) is not bool:
            raise MealError('跳过标记无效')
        data = {**row.data, 'skip':edit['skip']}
        if not edit['skip']:
            if row.data.get('period_conflict'):
                if edit.get('period_confirmed') is not True:
                    raise MealError('请先显式核对并确认文件/页内年月差异')
                data['period_confirmed'] = True
                data['correction_reason'] = required_text(edit.get('correction_reason'), '年月归属核对说明')
            amount = cents(edit.get('amount')) / 100
            department = required_text(edit.get('dept_name'), '部门', 100)
            if amount != row.data['original_amount'] or department != row.data['original_dept_name']:
                data['correction_reason'] = required_text(edit.get('correction_reason'), '更正说明')
            if row.data['kind'] == 'person':
                emp_id = edit.get('emp_id')
                emp = db.session.get(Employee, emp_id) if type(emp_id) is int else None
                if not emp:
                    raise MealError('请匹配现有人员或明确跳过')
                if emp.id in seen:
                    raise MealError('与已有原账或当前文件人员重复，请跳过冲突行', 409)
                if emp.id != row.emp_id:
                    data['correction_reason'] = required_text(edit.get('correction_reason'), '人员映射说明')
                seen.add(emp.id)
                row.emp_id = emp.id
            data.update(amount=amount, dept_name=department, error='', reviewed_by=operator)
        row.data = data
    record.status = 'confirmed'
    db.session.flush()
    return record
