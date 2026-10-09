"""Preview and atomically apply signed meal adjustment spreadsheets."""
from collections import Counter
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
import hashlib

from flask import current_app
from itsdangerous import BadSignature, URLSafeTimedSerializer
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from models import db
from models.meal_ticket import MealTicketBatch, MealTicketItem
from services.meal_ticket_import_service import identifier, label
from services.meal_ticket_service import MealError, adjustment, batch_for_write, cents, required_text, totals

HEADERS = ('工号', '姓名', '调整金额', '原因')


def serializer():
    return URLSafeTimedSerializer(current_app.config['SECRET_KEY'], salt='meal-adjustment-import-v1')


def batch_items(batch):
    return MealTicketItem.query.filter_by(batch_key=batch.key).all()


def already_imported(items, checksum):
    return any(receipt['file_digest'] == checksum for item in items
        for receipt in item.source.get('adjustment_imports', []))


def validate_rows(rows, items):
    people = {item.emp_no_snapshot:item for item in items}
    counts = Counter(row['emp_no'] for row in rows)
    result = []
    for row in rows:
        errors = []
        item = people.get(row['emp_no'])
        if not item:
            errors.append('工号未匹配本月核算人员')
        elif row['name'] != item.name:
            errors.append('姓名与本月核算人员不一致')
        if counts[row['emp_no']] > 1:
            errors.append('同一工号重复，请合并为一行')
        try:
            reason = required_text(row['reason'], '原因')
        except MealError as exc:
            errors.append(str(exc))
            reason = row['reason']
        amount = None
        try:
            if isinstance(row['amount'], str) and row['amount'].startswith('='):
                raise MealError('调整金额不能使用公式，请填写实际金额')
            value = cents(row['amount'])
            if value == 0:
                raise MealError('调整金额不能为零')
            if item and totals(item)[0] + value < 0:
                raise MealError('调整后应发金额不能为负数')
            amount = value / 100
        except MealError as exc:
            errors.append(str(exc))
        result.append({**row, 'reason':reason, 'amount':amount,
            'dept_name':item.dept_name if item else '', 'error':'；'.join(errors)})
    return result


def preview_adjustments(payload, filename, batch):
    if batch.status != 'confirmed':
        raise MealError('请先完成月度核算，再导入后续补扣', 409)
    if not filename.lower().endswith('.xlsx') or not payload or len(payload) > 20 * 1024 * 1024:
        raise MealError('请选择不超过 20 MB 的 xlsx 文件')
    checksum = hashlib.sha256(payload).hexdigest()
    items = batch_items(batch)
    if already_imported(items, checksum):
        raise MealError('该文件已导入本月补扣，不能重复入账', 409)
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise MealError('工作簿解压内容过大')
        book = load_workbook(BytesIO(payload), read_only=True, data_only=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_row and sheet.max_row > 10001:
                raise MealError('补扣清单不能超过 10000 行')
            if sheet.max_column and sheet.max_column > 32:
                raise MealError('工作簿列数过多，请使用不超过 32 列的导入模板')
            iterator = sheet.iter_rows(max_row=sheet.max_row or 10002, max_col=32, values_only=True)
            headers = [label(value) for value in next(iterator, ())]
            if any(headers.count(header) != 1 for header in HEADERS):
                raise MealError('请使用导入模板，必须包含工号、姓名、调整金额、原因且表头不能重复')
            indices = [headers.index(header) for header in HEADERS]
            rows = []
            for number, values in enumerate(iterator, 2):
                if number > 10001:
                    raise MealError('补扣清单不能超过 10000 行')
                cells = [values[index] if index < len(values) else None for index in indices]
                if all(value is None or str(value).strip() == '' for value in cells):
                    continue
                emp_no, name, amount, reason = cells
                rows.append({'row':number, 'emp_no':identifier(emp_no), 'name':str(name or '').strip(),
                    'amount':amount, 'reason':str(reason or '').strip()})
        finally:
            book.close()
    except MealError:
        raise
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError, ParseError):
        raise MealError('无法读取 Excel 文件，请使用有效的 xlsx 导入模板')
    if not rows:
        raise MealError('补扣清单没有可导入的数据')
    validated = validate_rows(rows, items)
    filename = Path(filename).name[:255]
    token = serializer().dumps({'batch_id':batch.id, 'batch_key':batch.key, 'version':batch.version,
        'filename':filename, 'file_digest':checksum, 'rows':validated})
    return {'batch_id':batch.id, 'version':batch.version, 'recharge_month':batch.recharge_month,
        'filename':filename, 'rows':validated, 'token':token,
        'error_count':sum(bool(row['error']) for row in validated),
        'total_amount':sum(cents(row['amount']) for row in validated if row['amount'] is not None) / 100}


def apply_adjustments(token, operator):
    if not isinstance(token, str):
        raise MealError('请先上传并预览补扣清单')
    try:
        preview = serializer().loads(token, max_age=7200)
    except BadSignature:
        raise MealError('导入预览无效或已过期，请重新上传预览')
    batch = MealTicketBatch.query.filter_by(id=preview['batch_id']).with_for_update().first()
    if not batch or batch.key != preview['batch_key'] or batch.status != 'confirmed':
        raise MealError('核算状态已变化，请重新上传预览', 409)
    items = batch_items(batch)
    if already_imported(items, preview['file_digest']):
        return batch
    batch_for_write(batch.id, preview['version'])
    if any(row['error'] for row in preview['rows']):
        raise MealError('清单存在错误，请修正 Excel 后重新上传预览')
    rows = validate_rows(preview['rows'], items)
    if any(row['error'] for row in rows):
        raise MealError('清单校验失败，请重新上传预览')
    people = {item.emp_no_snapshot:item for item in items}
    for row in rows:
        item = people[row['emp_no']]
        adjustment(batch, item.id, row['amount'], row['reason'], operator)
        item.source = {**item.source, 'adjustment_imports':[*item.source.get('adjustment_imports', []),
            {'file_digest':preview['file_digest'], 'filename':preview['filename'], 'operator':operator,
             'amount_cents':cents(row['amount']), 'reason':row['reason'], 'created_at':datetime.utcnow().isoformat()}]}
    db.session.flush()
    return batch
