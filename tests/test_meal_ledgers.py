import unittest
import io
from openpyxl import Workbook, load_workbook

from tests import test_meal_tickets as meal_tests


class MealLedgerTests(unittest.TestCase):
    setUp = meal_tests.MealTicketTests.setUp
    tearDown = meal_tests.MealTicketTests.tearDown
    generate = meal_tests.MealTicketTests.generate
    confirm = meal_tests.MealTicketTests.confirm
    post = meal_tests.MealTicketTests.post

    def ledger(self, kind, **values):
        return self.client.post('/api/meal-ledgers/' + kind, json={
            'month': '2026-09', 'date': '2026-09-30', 'request_key': kind + '-1',
            **values}, headers=self.headers)

    def test_clearance_does_not_reopen_settled_subsidy(self):
        batch = self.confirm(self.generate())
        paid = self.post('/payments', {'batch_id': batch['id'], 'version': batch['version'],
            'item_id': batch['items'][0]['id'], 'amount': '176', 'date': '2026-09-01',
            'reference': '充值凭证', 'request_key': 'paid-1', 'kind': 'recharge'})
        self.assertEqual(paid.status_code, 200, paid.get_json())
        response = self.ledger('clearance', emp_no='001', name='员工甲', dept_name='未分配部门', amount='20')
        self.assertEqual(response.status_code, 200, response.get_json())
        data = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(data['items'][0]['difference'], 0)
        year = self.client.get('/api/meal-ledgers/annual?year=2026', headers=self.headers).get_json()
        self.assertEqual(year['months'][8]['recovered_amount'], 20)
        self.assertEqual(year['months'][8]['recharge_amount'], 176)
        self.assertIsNone(year['months'][8]['consumption_amount'])

    def test_external_categories_and_zero_consumption_are_distinct(self):
        self.assertEqual(self.ledger('external', name='访客', dept_name='行政部', category='card', amount='50').status_code, 200)
        self.assertEqual(self.ledger('external', name='访客', dept_name='行政部', category='paper', amount='30', request_key='paper-1').status_code, 200)
        self.assertEqual(self.ledger('consumption', floor2='0', floor3='12.50').status_code, 200)
        data = self.client.get('/api/meal-ledgers/annual?year=2026', headers=self.headers).get_json()
        self.assertEqual(data['months'][8]['recharge_amount'], 50)
        self.assertEqual(data['months'][8]['paper_amount'], 30)
        self.assertEqual(data['months'][8]['consumption_amount'], 12.5)
        self.assertIsNone(data['months'][7]['consumption_amount'])

    def test_department_includes_guest_card_and_paper_once_and_excludes_voids(self):
        batch = self.confirm(self.generate())
        self.post('/payments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':176, 'kind':'recharge',
            'date':'2026-09-01', 'reference':'员工充值', 'request_key':'dept-employee'})
        values = dict(name='客人', dept_name='未分配部门', category='card', amount='200')
        self.assertEqual(self.ledger('external', **values).status_code, 200)
        self.assertEqual(self.ledger('external', **values).status_code, 200)
        paper = self.ledger('external', name='客人', dept_name='未分配部门', category='paper',
            amount='100', request_key='dept-paper').get_json()
        self.ledger('external', name='客人乙', dept_name='行政部', category='card',
            amount='50', request_key='dept-other')
        self.ledger('external', name='次月客人', dept_name='行政部', category='paper', amount='99',
            month='2026-10', date='2026-10-01', request_key='dept-next')
        data = self.client.get('/api/meal-ledgers/department?month=2026-09', headers=self.headers).get_json()
        rows = {r['dept_name']:r for r in data['items']}
        self.assertEqual(rows['未分配部门']['paid_amount'], 176)
        self.assertEqual(rows['未分配部门']['external_card_amount'], 200)
        self.assertEqual(rows['未分配部门']['external_paper_amount'], 100)
        self.assertEqual(rows['未分配部门']['total_paid_amount'], 476)
        self.assertEqual(rows['行政部']['total_paid_amount'], 50)
        export = self.client.get('/api/meal-ledgers/export?report=department&month=2026-09', headers=self.headers)
        book = load_workbook(io.BytesIO(export.data))
        self.assertEqual(book.active.cell(book.active.max_row, 2).value, 526)
        detail = book['09月明细']
        self.assertEqual([detail.cell(3, n).value for n in range(1, 6)],
            ['部门', '员工净实际发放', '客人卡充值', '客人纸质菜票', '部门发放合计'])
        self.assertEqual(detail.cell(detail.max_row, 5).value, 526)
        restricted = self.client.get('/api/meal-ledgers/export?report=department&month=2026-09',
            headers={'Authorization':'Bearer ' + self.viewer_token})
        self.assertEqual(restricted.status_code, 200)
        restricted_book = load_workbook(io.BytesIO(restricted.data))
        restricted_detail = restricted_book['09月明细']
        self.assertEqual(list(restricted_detail.values)[3], ('未分配部门', 176, 0, 0, 176))
        self.assertEqual(restricted_detail.max_row, 5)
        self.client.post(f'/api/meal-ledgers/records/{paper["id"]}/void',
            json={'reason':'重复纸票'}, headers=self.headers)
        data = self.client.get('/api/meal-ledgers/department?month=2026-09', headers=self.headers).get_json()
        row = next(r for r in data['items'] if r['dept_name'] == '未分配部门')
        self.assertEqual(row['total_paid_amount'], 376)
        from services.meal_ledger_service import departments
        with self.app.app_context():
            limited = departments('2026-09', [self.emp_id])
        self.assertEqual(len(limited['items']), 1)
        self.assertEqual(limited['items'][0]['total_paid_amount'], 176)

    def test_guest_only_department_is_visible_without_employee_batch_and_requires_department(self):
        response = self.ledger('external', name='客人', category='paper', amount='100')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.ledger('external', name='客人', dept_name='接待部', category='paper', amount='100').status_code, 200)
        data = self.client.get('/api/meal-ledgers/department?month=2026-09', headers=self.headers).get_json()
        self.assertEqual(data['items'][0]['dept_name'], '接待部')
        self.assertEqual(data['items'][0]['total_paid_amount'], 100)
        self.assertEqual(data['items'][0]['count'], 0)

    def test_retry_and_void_preserve_history_without_double_counting(self):
        values = dict(emp_no='guest', name='客人卡', amount='20')
        one = self.ledger('clearance', **values)
        self.assertEqual(one.status_code, 200)
        two = self.ledger('clearance', **values)
        self.assertEqual(one.get_json()['id'], two.get_json()['id'])
        self.assertEqual(self.ledger('clearance', **{**values, 'amount': '21'}).status_code, 409)
        response = self.client.post('/api/meal-ledgers/records/%s/void' % one.get_json()['id'],
            json={'reason': '重复录入'}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        rows = self.client.get('/api/meal-ledgers/clearance?month=2026-09', headers=self.headers).get_json()
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0]['voided'])
        annual = self.client.get('/api/meal-ledgers/annual?year=2026', headers=self.headers).get_json()
        self.assertEqual(annual['months'][8]['recovered_amount'], 0)

    def test_person_query_permission_cannot_read_whole_company_ledgers(self):
        for path in ('external?month=2026-09', 'clearance?month=2026-09', 'annual?year=2026'):
            response = self.client.get('/api/meal-ledgers/' + path,
                headers={'Authorization': 'Bearer ' + self.viewer_token})
            self.assertEqual(response.status_code, 403)

    def test_invalid_amount_period_and_category_are_rejected(self):
        for change in ({'amount': '-1'}, {'amount': 'NaN'}, {'month': '2026-13'},
                       {'category': 'other'}, {'date': '2026-08-31'}):
            response = self.ledger('external', **{'name': '访客', 'category': 'card', 'amount': '50', **change})
            self.assertEqual(response.status_code, 400, response.get_json())

    def test_clearance_import_preview_confirm_and_retry(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '8月'
        sheet.append(['人员编号', '人员姓名', '部门名称', '卡余额', '备注'])
        sheet.append(['0001', '客人用', '客人', 20.5, '清零'])
        output = io.BytesIO()
        book.save(output)
        response = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'clearance', 'month': '2026-08', 'file': (io.BytesIO(output.getvalue()), '取款.xlsx')},
            headers=self.headers)
        self.assertEqual(response.status_code, 200, response.get_json())
        preview = response.get_json()
        self.assertEqual(self.client.get('/api/meal-ledgers/clearance?month=2026-08', headers=self.headers).get_json(), [])
        self.assertEqual(preview['rows'][0]['emp_no'], '0001')
        response = self.client.post('/api/meal-ledgers/imports/%s/confirm' % preview['id'],
            json={'rows': [{'index': 0, 'month': '2026-08', 'date': '2026-08-31'}]}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.get_json())
        response = self.client.post('/api/meal-ledgers/imports/%s/confirm' % preview['id'],
            json={'rows': [{'index': 0, 'month': '2026-08', 'date': '2026-08-31'}]}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        rows = self.client.get('/api/meal-ledgers/clearance?month=2026-08', headers=self.headers).get_json()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['amount'], 20.5)
        response = self.client.get('/api/meal-ledgers/export?report=clearance&month=2026-08', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        saved = load_workbook(io.BytesIO(response.data))
        self.assertEqual(saved.active['A4'].value, '0001')
        self.assertEqual(saved.active['E4'].value, 20.5)

    def test_clearance_preview_preserves_processing_date(self):
        book = Workbook(); sheet = book.active
        sheet.append(['人员编号', '人员姓名', '卡余额', '处理日期'])
        sheet.append(['0001', '客人卡', 20, '2026-08-12'])
        raw = io.BytesIO(); book.save(raw)
        response = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'clearance', 'month': '2026-08', 'file': (io.BytesIO(raw.getvalue()), '取款.xlsx')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()['rows'][0]['date'], '2026-08-12')

    def test_archive_recharge_uses_paid_amount_and_authorized_scope(self):
        self.confirm(self.generate())
        response = self.client.get('/api/meal-ledgers/export?report=recharge&month=2026-09',
            headers={'Authorization': 'Bearer ' + self.viewer_token})
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['员工异常查询1', '员工充值记录', '管理人员查询'])
        sheet = book['员工充值记录']
        self.assertEqual(sheet['B4'].value, '001')
        self.assertEqual(sheet['J4'].value, 0)
        self.assertEqual(sheet.max_row, 4)

    def test_department_archive_keeps_two_blocks_and_paid_total(self):
        self.confirm(self.generate())
        for index, name in enumerate(('行政部', '财务部')):
            self.ledger('department', dept_name=name, registrar='登记人', request_key='department-' + str(index))
        response = self.client.get('/api/meal-ledgers/export?report=department&month=2026-09', headers=self.headers)
        book = load_workbook(io.BytesIO(response.data))
        sheet = book.active
        self.assertEqual([sheet.cell(3, n).value for n in (1, 2, 3, 4, 6, 7, 8, 9)],
                         ['部门', '金额', '登记人', '备注'] * 2)
        names = [sheet.cell(row, col).value for row in range(4, sheet.max_row) for col in (1, 6)]
        self.assertIn('行政部', names)
        self.assertIn('财务部', names)
        self.assertEqual(sheet.cell(sheet.max_row, 1).value, '合计')
        self.assertEqual(sheet.cell(sheet.max_row, 2).value, 0)

    def test_navigation_separates_query_operations_and_downloads(self):
        modules = self.client.get('/api/query/navigation', headers=self.headers).get_json()['modules']
        by_slug = {m['slug']: m for m in modules}
        query_paths = [r['href'] for r in by_slug['query']['entries']]
        self.assertIn('/employee/meal-ticket-query', query_paths)
        self.assertNotIn('/employee/summary-download', query_paths)
        self.assertEqual([r['href'] for r in by_slug['downloads']['entries']], ['/downloads/attendance', '/downloads/meal-tickets'])

    def test_business_month_backup_contains_clearance_without_attendance_link(self):
        from models.account_set import AccountSet
        from services.account_set_backup_service import export_backup, read_backup
        response = self.ledger('clearance', month='2026-08', date='2026-08-31', name='客人卡', amount='20')
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            account = AccountSet.query.filter_by(month='2026-08').first()
            archive = export_backup(account.id)
            document = read_backup(archive.getvalue() if hasattr(archive, 'getvalue') else archive)
        self.assertEqual(document['datasets']['meal_ledger_records'][0]['amount_cents'], 2000)

    def test_import_source_and_business_month_records_restore_together(self):
        from pathlib import Path
        from models import db
        from models.account_set import AccountSet
        from models.meal_ledger import MealLedgerRecord, MealLedgerImport
        from services.account_set_backup_service import export_backup, read_backup
        from services.account_set_restore_service import build_preview, restore_backup
        book = Workbook(); sheet = book.active; sheet.title = '8月'
        sheet.append(['人员编号', '人员姓名', '卡余额'])
        sheet.append(['0001', '客人卡', 20])
        raw = io.BytesIO(); book.save(raw)
        preview = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'clearance', 'month': '2026-09', 'file': (io.BytesIO(raw.getvalue()), '取款.xlsx')}, headers=self.headers).get_json()
        response = self.client.post('/api/meal-ledgers/imports/%s/confirm' % preview['id'],
            json={'rows': [{'index': 0, 'month': '2026-08', 'date': '2026-08-31'}]}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.get_json())
        with self.app.app_context():
            account = AccountSet.query.filter_by(month='2026-08').one()
            account.is_locked = False
            db.session.commit()
            document = read_backup(export_backup(account.id))
            self.assertEqual(len(document['datasets']['meal_ledger_imports']), 1)
            MealLedgerRecord.query.delete(); MealLedgerImport.query.delete()
            db.session.commit()
            restored = build_preview(document, {})
            self.assertEqual(restored['blockers'], [])
            restore_backup(document, {}, {}, restored['fingerprint'], self.admin_id)
            self.assertEqual(MealLedgerRecord.query.one().month, '2026-08')
            self.assertEqual(MealLedgerImport.query.one().month, '2026-09')
            self.assertEqual(Path(MealLedgerImport.query.one().stored_path).read_bytes(), raw.getvalue())
            again = build_preview(document, {})
            restore_backup(document, {}, {}, again['fingerprint'], self.admin_id)
            self.assertEqual(MealLedgerRecord.query.count(), 1)

    def test_imported_consumption_keeps_historical_totals_as_comparison(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '全年菜票录入汇总'
        sheet.append(['月份', '充值金额', '消费金额', '二楼消费金额', '三楼消费金额', '收回金额'])
        sheet.append(['8月', 1000, 750, 500, 250, 250])
        raw = io.BytesIO(); book.save(raw)
        response = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'consumption', 'month': '2026-08', 'file': (io.BytesIO(raw.getvalue()), '全年.xlsx')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        data = response.get_json()
        response = self.client.post('/api/meal-ledgers/imports/%s/confirm' % data['id'], json={
            'rows': [{'index': 0, 'month': '2026-08', 'date': '2026-08-31'}]}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        data = self.client.get('/api/meal-ledgers/annual?year=2026', headers=self.headers).get_json()
        self.assertEqual(data['months'][7]['recharge_amount'], 0)
        self.assertEqual(data['months'][7]['recovered_amount'], 0)
        self.assertEqual(data['months'][7]['historical_recharge_amount'], 1000)
        self.assertEqual(data['months'][7]['historical_recovered_amount'], 250)

    def test_yearly_clearance_archive_has_month_tabs(self):
        self.ledger('clearance', name='客人', amount='20')
        response = self.client.get('/api/meal-ledgers/export?report=clearance&year=2026', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data))
        self.assertIn('09月', book.sheetnames)
        self.assertEqual(book['09月']['E4'].value, 20)

    def test_query_detail_shows_clearance_separately_from_payments(self):
        self.confirm(self.generate())
        self.ledger('clearance', emp_no='001', name='员工甲', amount='20')
        response = self.client.get('/api/meal-tickets?recharge_month=2026-09',
            headers={'Authorization': 'Bearer ' + self.viewer_token})
        data = response.get_json()
        self.assertEqual(len(data['items']), 1)
        self.assertEqual(data['items'][0]['clearances'][0]['amount'], 20)
        self.assertEqual(data['items'][0]['paid_amount'], 0)

    def test_note_update_preserves_original_department_import_amount(self):
        self.assertEqual(self.ledger('department', dept_name='生产部', historical_amount='100').status_code, 200)
        self.assertEqual(self.ledger('department', dept_name='生产部', registrar='登记人', request_key='note-2').status_code, 200)
        data = self.client.get('/api/meal-ledgers/department?month=2026-09', headers=self.headers).get_json()
        self.assertEqual(data['items'][0]['historical_amount'], 100)

    def test_restore_preview_blocks_active_month_slot_conflict(self):
        from models.account_set import AccountSet
        from services.account_set_backup_service import export_backup, read_backup
        from services.account_set_restore_service import build_preview, validate_selection
        from models import db
        from models.meal_ledger import MealLedgerRecord
        self.ledger('consumption', month='2026-08', date='2026-08-31', floor2='10', floor3='20')
        with self.app.app_context():
            account = AccountSet.query.filter_by(month='2026-08').first()
            doc = read_backup(export_backup(account.id))
            old = MealLedgerRecord.query.one()
            old.key = 'changed-source-key'
            old.request_key = 'changed-request-key'
            db.session.commit()
            preview = build_preview(doc, {})
            blockers = validate_selection(doc, {}, preview['rows'], {})
            self.assertTrue(any(b['code'] == 'meal_ledger_unique_conflict' for b in blockers))

    def test_scoped_query_does_not_expose_other_person_pending_refunds(self):
        from models.meal_ticket import MealTicketBatch
        from models import db
        self.confirm(self.generate())
        with self.app.app_context():
            batch = MealTicketBatch.query.one()
            batch.reconciliation = {'pending_refunds': [{'emp_no': '002', 'name': '员工乙', 'amount': 99}]}
            db.session.commit()
        data = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers={'Authorization': 'Bearer ' + self.viewer_token}).get_json()
        self.assertIsNone(data['reconciliation'])

    def test_two_previews_cannot_import_same_clearance_without_duplicate_confirmation(self):
        identifiers = []
        for title in ('8月', '8月原账副本'):
            book = Workbook(); sheet = book.active; sheet.title = title
            sheet.append(['人员编号', '人员姓名', '卡余额'])
            sheet.append(['001', '员工甲', 20])
            raw = io.BytesIO(); book.save(raw)
            response = self.client.post('/api/meal-ledgers/imports/preview', data={
                'kind': 'clearance', 'month': '2026-08', 'file': (io.BytesIO(raw.getvalue()), '取款.xlsx')}, headers=self.headers)
            self.assertEqual(response.status_code, 200)
            identifiers.append(response.get_json()['id'])
        choices = {'rows': [{'index': 0, 'month': '2026-08', 'date': '2026-08-31'}]}
        self.assertEqual(self.client.post('/api/meal-ledgers/imports/%s/confirm' % identifiers[0], json=choices, headers=self.headers).status_code, 200)
        self.assertEqual(self.client.post('/api/meal-ledgers/imports/%s/confirm' % identifiers[1], json=choices, headers=self.headers).status_code, 400)
        choices['rows'][0]['accept_duplicate'] = True
        self.assertEqual(self.client.post('/api/meal-ledgers/imports/%s/confirm' % identifiers[1], json=choices, headers=self.headers).status_code, 200)

    def test_clearance_import_excludes_total_rows(self):
        book = Workbook(); sheet = book.active
        sheet.append(['人员编号', '人员姓名', '卡余额'])
        sheet.append(['001', '员工甲', 20])
        sheet.append([None, '合计', 20])
        raw = io.BytesIO(); book.save(raw)
        response = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'clearance', 'month': '2026-08', 'file': (io.BytesIO(raw.getvalue()), '取款.xlsx')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.get_json()['rows']), 1)

    def test_external_import_ignores_right_hand_pivot_totals(self):
        book = Workbook(); sheet = book.active
        sheet.append(['2025年外来人员菜票领用', None, None, None, None, None, None, None, '=SUM(I3:I55)', None, None, None, None, '月份', '部门', '类别', '发放金额（元）'])
        sheet.append(['月份', '日期', '填表人', '部门', '外来人员（部门）', '姓名', '时间', '天数', '发放金额（元）', '类别', '卡号', None, None, '月份', '部门', '类别', '发放金额（元）'])
        sheet.append(['8月', None, '登记人', '行政部', '外来单位', '访客甲', '', 1, 20, '充卡', '001', None, None, '总计', '其他部门', '纸质', 999])
        raw = io.BytesIO(); book.save(raw)
        response = self.client.post('/api/meal-ledgers/imports/preview', data={
            'kind': 'external', 'month': '2026-08', 'file': (io.BytesIO(raw.getvalue()), '外来.xlsx')}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        row = response.get_json()['rows'][0]
        self.assertEqual(row['amount'], 20)
        self.assertEqual(row['dept_name'], '行政部')
        self.assertEqual(row['category'], 'card')
        self.assertIn('年份', row['warning'])

    def test_payment_restore_cannot_reuse_clearance_database_source(self):
        from datetime import date
        from models import db
        from models.account_set import AccountSet
        from models.meal_ticket import MealTicketBatch, MealTicketItem, MealTicketPayment
        from services.account_set_backup_service import export_backup, read_backup
        from services.account_set_restore_service import build_preview
        from services.meal_ledger_service import create
        self.confirm(self.generate())
        with self.app.app_context():
            batch = MealTicketBatch.query.one()
            item = MealTicketItem.query.filter_by(batch_key=batch.key).first()
            payment = MealTicketPayment(item_key=item.key, month=batch.month, kind='refund', amount_cents=-2000,
                payment_date=date(2026, 9, 30), reference='原数据库扣回', operator='admin',
                request_key='card-meal:test:refund:1', request_digest='x' * 64)
            db.session.add(payment); db.session.commit()
            doc = read_backup(export_backup(batch.account_set_id))
            db.session.delete(payment)
            AccountSet.query.one().is_locked = False
            create('clearance', {'month':'2026-09', 'date':'2026-09-30', 'emp_no':'001', 'name':'员工甲', 'amount':'20', 'request_key':'cleared-db'}, 'admin', source_key='card-meal:test:refund:1')
            db.session.commit()
            preview = build_preview(doc, {})
            self.assertTrue(any(b['code'] == 'meal_ledger_unique_conflict' for b in preview['blockers']))
