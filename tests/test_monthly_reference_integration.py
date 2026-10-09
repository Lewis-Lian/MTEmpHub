"""Historical consumers must not read mutable current reference attributes."""
from datetime import date
from io import BytesIO

import openpyxl
from flask import g
import pytest

from models import db
from models.department import Department
from models.employee import Employee
from models.daily_record import DailyRecord
from models.monthly_report import MonthlyReport
from models.employee_shift import EmployeeShiftAssignment
from models.shift import Shift
from models.user import User, UserDepartmentAssignment
from services.monthly_reference_service import ensure_month_reference
from tests.test_account_set_backup import backup_app
import routes.query_core  # Register consumer models before fixture create_all.


@pytest.fixture
def historical_months(backup_app):
    employee = Employee.query.one()
    DailyRecord.query.delete()
    shift = Shift(shift_no='S1', shift_name='白班', time_slots=[['08:00', '12:00'], ['13:00', '17:00']])
    warehouse = Department(dept_no='D2', dept_name='仓储')
    manager = Employee(emp_no='M1', name='经理', dept_id=employee.dept_id, is_manager=True)
    db.session.add_all([shift, warehouse, manager])
    db.session.flush()
    db.session.add(EmployeeShiftAssignment(emp_id=employee.id, shift_id=shift.id))
    for month in ('2026-06', '2026-07'):
        if month.endswith('07'):
            employee.dept_id = warehouse.id
        ensure_month_reference(month)
        db.session.add(DailyRecord(emp_id=employee.id, record_date=date.fromisoformat(month + '-04'),
            shift_id=shift.id, employee_payload={'actual_hours': 8, 'expected_hours': 8,
                'check_in_times': ['08:00'], 'check_out_times': ['17:00'],
                'raw_data': {'刷卡时间数据': '08:00,17:00'}},
            manager_payload={'actual_hours': 120, 'raw_data': {'上班1打卡时间': '08:00'}}))
        db.session.add(MonthlyReport(emp_id=manager.id, report_month=month,
                                    manager_raw_data={'出勤天数': 20}))
    db.session.commit()
    yield employee, manager, shift


def change_current(employee, manager, shift):
    employee.name = '现名'
    employee.dept_id = None
    employee.is_manager = True
    employee.employee_stats_attendance_source = 'manager'
    employee.resigned_at = date(2026, 5, 1)
    employee.is_active = False
    manager.name = '现经理'
    manager.is_manager = False
    manager.resigned_at = date(2026, 5, 1)
    manager.is_active = False
    Department.query.filter_by(dept_no='D1').one().dept_name = '现部门'
    shift.time_slots = [['08:00', '17:00']]
    db.session.commit()


@pytest.mark.parametrize('month,department', [('2026-06', '生产'), ('2026-07', '仓储')])
def test_query_calendar_summary_and_meal_recalculation_use_month_identity(historical_months, backup_app, month, department):
    from routes.query_core import _build_attendance_calendar_payload, _build_final_rows, _build_department_hours_rows
    from services.attendance_service import AttendanceService
    from services.meal_ticket_service import source_snapshot
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    with backup_app.test_request_context():
        calendar = _build_attendance_calendar_payload(employee, month)
        assert calendar['employee']['name'] == '甲'
        assert calendar['employee']['dept_name'] == department
        assert _build_final_rows(month, [employee.id])[0][:4] == [department, 'E1', '甲', 1.0]
        assert _build_department_hours_rows(month, [employee.id]) == [
            {'dept_name': department, 'total_hours': 8.0, 'member_count': 1}]
        assert AttendanceService.monthly_summary(employee.id, month)['actual_hours'] == 8
        meal = next(row for row in source_snapshot(month) if row['emp_id'] == employee.id)
        assert (meal['name'], meal['dept_name'], meal['is_manager'], meal['days']) == ('甲', department, False, 1)


def test_manager_membership_and_recalculation_survive_current_type_change(historical_months):
    from services.manager_attendance_service import build_manager_rows, ManagerAttendanceOptions
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    rows = build_manager_rows(ManagerAttendanceOptions(month='2026-06'), sync_month_stats=True)
    assert [(r['emp_no'], r['name'], r['dept_name']) for r in rows] == [('M1', '经理', '生产')]
    assert rows[0]['attendance_days'] == 20


def test_permissions_use_current_grants_on_historical_department_membership(historical_months, backup_app):
    from routes.query_core import _accessible_emp_ids, _non_manager_emp_ids, departments_api
    employee, manager, shift = historical_months
    department = Department.query.filter_by(dept_no='D1').one()
    viewer = User(username='viewer', role='readonly', password_hash='unused', page_permissions={'employee_dashboard': True})
    db.session.add(viewer)
    db.session.flush()
    grant = UserDepartmentAssignment(user_id=viewer.id, dept_id=department.id)
    db.session.add(grant)
    db.session.commit()
    change_current(employee, manager, shift)
    with backup_app.test_request_context('/?month=2026-06'):
        g.current_user = viewer
        assert employee.id in _non_manager_emp_ids(_accessible_emp_ids())
        assert departments_api().get_json()[0]['dept_name'] == '生产'
    with backup_app.test_request_context('/?month=2026-07'):
        g.current_user = viewer
        assert employee.id not in _accessible_emp_ids()
    db.session.delete(grant)
    db.session.commit()
    with backup_app.test_request_context('/?month=2026-06'):
        g.current_user = viewer
        assert _accessible_emp_ids() == []


def test_history_download_contains_frozen_name_department_and_hours(historical_months, backup_app):
    from routes.query_core import punch_records_export_api
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    with backup_app.test_request_context('/?month=2026-07&emp_id=' + str(employee.id)):
        g.current_user = User.query.filter_by(username='admin').one()
        response = punch_records_export_api()
        assert not isinstance(response, tuple), 'historical employee must remain downloadable'
        response.direct_passthrough = False
        workbook = openpyxl.load_workbook(BytesIO(response.get_data()))
        row = list(workbook.active.values)[1]
        assert row[1:4] == ('E1', '甲', '仓储')
        assert row[8] == 8


def test_legacy_calendar_reports_missing_history_without_writing(backup_app):
    from routes.query_core import _build_attendance_calendar_payload
    from models.monthly_reference_snapshot import MonthlyReferenceSnapshot
    with backup_app.test_request_context():
        payload = _build_attendance_calendar_payload(Employee.query.one(), '2026-06')
        assert payload['reference']['quality'] == 'missing'
        assert payload['reference']['warnings']
        assert MonthlyReferenceSnapshot.query.count() == 0


def test_daily_import_captures_month_and_preserves_previous_month(backup_app):
    from services.import_service import ImportService
    from services.monthly_reference_service import get_month_reference
    rows = [['人员编号', '考勤日期', '实出勤小时', '刷卡时间数据'],
            ['E1', '2026-07-04', 8, '08:00,17:00']]
    assert ImportService._import_daily_records(rows)['imported'] == 1
    assert get_month_reference('2026-07', 'employee', 'E1')['name'] == '甲'
    Employee.query.one().name = '乙'
    db.session.commit()
    ImportService._import_daily_records(rows)
    assert get_month_reference('2026-07', 'employee', 'E1')['name'] == '甲'
    assert get_month_reference('2026-06', 'employee', 'E1') is None


def test_current_maintenance_excludes_inactive_references(backup_app):
    from routes.admin_core import employees_list, departments_list, list_shifts
    shift = Shift(shift_no='S1', shift_name='退出班次', time_slots=[], is_active=False)
    db.session.add(shift)
    Employee.query.one().is_active = False
    Department.query.one().is_active = False
    db.session.commit()
    with backup_app.test_request_context():
        g.current_user = User.query.filter_by(username='admin').one()
        assert employees_list().get_json() == []
        assert departments_list().get_json() == []
        assert list_shifts().get_json() == []


def test_reference_change_requires_all_affected_month_categories(historical_months):
    from services.monthly_reference_service import reference_change_blockers
    from models.meal_ticket import MealTicketBatch, MealTicketItem
    employee, _, _ = historical_months
    batch = MealTicketBatch(key='B1', month='2026-06', recharge_month='2026-07', account_set_id=1,
                            source_digest='digest', created_by='admin')
    db.session.add(batch)
    db.session.flush()
    db.session.add(MealTicketItem(key='I1', batch_key='B1', month='2026-06', emp_id=employee.id,
                                 emp_no_snapshot='E1', name='甲', dept_name='生产', source={},
                                 is_manager=False, days=1, base_cents=800))
    db.session.commit()
    proposed = [{'kind': 'employee', 'business_key': 'E1', 'payload': {'name': '修正'}}]
    blockers = reference_change_blockers('2026-06', proposed, ['attendance'])
    assert blockers[0]['required_categories'] == ['attendance', 'meal_tickets']
    assert blockers[0]['month'] == '2026-06'
    assert reference_change_blockers('2026-06', proposed, ['attendance', 'meal_tickets']) == []
    assert reference_change_blockers('2026-07', proposed, ['attendance']) == []


def test_explicit_month_correction_changes_only_selected_month(historical_months):
    from services.monthly_reference_service import correct_month_reference, get_month_reference
    from models.account_set import AccountSet
    before = get_month_reference('2026-07', 'employee', 'E1')
    correct_month_reference('2026-06', 'employee', 'E1', {'name': '月内修正'}, ['attendance'])
    db.session.commit()
    assert get_month_reference('2026-06', 'employee', 'E1')['name'] == '月内修正'
    assert get_month_reference('2026-07', 'employee', 'E1') == before
    assert Employee.query.filter_by(emp_no='E1').one().name == '甲'
    AccountSet.query.one().is_locked = True
    db.session.commit()
    with pytest.raises(ValueError, match='锁定'):
        correct_month_reference('2026-06', 'employee', 'E1', {'name': '非法修正'}, ['attendance'])


def test_legacy_evidence_requires_explicit_choice_and_keeps_partial_quality(backup_app):
    from services.monthly_reference_service import reconcile_month_evidence, get_month_reference
    from models.monthly_reference_snapshot import MonthlyReferenceSnapshot
    DailyRecord.query.one().raw_data = {'姓名': '历史甲', '部门': '生产'}
    db.session.commit()
    ensure_month_reference('2026-06')
    reconcile_month_evidence('2026-06', [{'business_key': 'E1', 'field': 'name', 'value': '历史甲'}], ['attendance'])
    db.session.commit()
    assert get_month_reference('2026-06', 'employee', 'E1')['name'] == '历史甲'
    snapshot = MonthlyReferenceSnapshot.query.filter_by(month='2026-06', kind='employee', business_key='E1').one()
    assert snapshot.quality == 'partial'
    assert snapshot.provenance['fields']['name']['sources'][0]['table'] == 'daily_records'
    assert snapshot.provenance['fields']['is_manager'] == 'baseline'
    with pytest.raises(ValueError, match='证据'):
        reconcile_month_evidence('2026-06', [{'business_key': 'E1', 'field': 'name', 'value': '无证据'}], ['attendance'])


def test_late_offset_uses_monthly_manager_identity(historical_months):
    from services.late_offset_service import late_offset_candidates
    employee, manager, shift = historical_months
    db.session.add(DailyRecord(emp_id=manager.id, record_date=date(2026, 6, 5),
        manager_payload={'late_minutes': 10, 'raw_data': {'上班1打卡时间': '08:10', '上班1打卡结果': '迟到'}}))
    db.session.commit()
    change_current(employee, manager, shift)
    rows = late_offset_candidates('2026-06')
    assert rows[0]['emp_name'] == '经理'
    assert rows[0]['late_minutes'] == 10


def test_saved_annual_manager_rows_use_selected_month_identity(historical_months, backup_app):
    from routes.admin_core import _manager_overtime_values
    from models.manager_month_stat import ManagerMonthStat
    employee, manager, shift = historical_months
    db.session.add(ManagerMonthStat(emp_id=manager.id, year=2026, stat_type='overtime', m6=2, remaining=2))
    db.session.commit()
    change_current(employee, manager, shift)
    with backup_app.test_request_context('/?month=2026-06'):
        values = _manager_overtime_values(2026, [manager.id])
        assert values['经理']['m6'] == 2
        assert values['经理']['dept_name'] == '生产'


def test_monthly_report_reimport_accepts_frozen_employee_type(historical_months):
    from services.import_service import ImportService
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    result = ImportService._import_monthly_report([
        ['人员编号', '人员名称', '部门编号', '部门名称', '出勤天数'], ['E1', '甲', 'D1', '生产', 22]],
        '2026_6月员工月报.xls')
    assert result['imported'] == 1
    assert MonthlyReport.query.filter_by(emp_id=employee.id, report_month='2026-06').one().employee_raw_data['出勤天数'] == 22


def test_month_correction_rejects_dangling_reference(historical_months):
    from services.monthly_reference_service import correct_month_reference, get_month_reference
    before = get_month_reference('2026-06', 'employee', 'E1')
    with pytest.raises(ValueError, match='部门'):
        correct_month_reference('2026-06', 'employee', 'E1', {'dept_no': 'MISSING'}, ['attendance'])
    assert get_month_reference('2026-06', 'employee', 'E1') == before


def test_existing_month_import_does_not_enroll_current_new_hire(historical_months):
    from services.import_service import ImportService
    from services.monthly_reference_service import get_month_reference
    db.session.add(Employee(emp_no='NEW', name='后来入职'))
    db.session.commit()
    ImportService._import_daily_records([['人员编号', '考勤日期', '实出勤小时'], ['E1', '2026-06-04', 8]])
    assert get_month_reference('2026-06', 'employee', 'NEW') is None


def test_manager_reimport_uses_frozen_name_and_type(historical_months):
    from services.import_service import ImportService
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    result = ImportService._import_manager_monthly_report([
        ['工号', '姓名', '出勤天数'], ['', '', ''], ['M1', '经理', 21]], '2026_6月管理人员月报.xls')
    assert result['imported'] == 1
    assert MonthlyReport.query.filter_by(emp_id=manager.id, report_month='2026-06').one().manager_raw_data['出勤天数'] == 21


def test_department_reimport_reactivates_original_row(backup_app):
    from routes.admin_imports import import_departments_xlsx
    department = Department.query.one()
    original_id = department.id
    department.is_active = False
    db.session.commit()
    workbook = openpyxl.Workbook()
    workbook.active.append(['部门编号', '部门名称', '上级部门编号'])
    workbook.active.append(['D1', '恢复部门', ''])
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    with backup_app.test_request_context('/', method='POST', data={'file': (data, '部门.xlsx')}):
        response = import_departments_xlsx.__wrapped__()
        assert response.get_json()['status'] == 'ok'
    assert Department.query.one().id == original_id
    assert Department.query.one().is_active is True


def test_current_maintenance_permissions_do_not_use_historical_department(historical_months, backup_app):
    from routes.admin_core import employees_list
    employee, _, _ = historical_months
    employee.dept_id = Department.query.filter_by(dept_no='D2').one().id
    viewer = User(username='current-viewer', role='readonly', password_hash='unused')
    db.session.add(viewer)
    db.session.flush()
    db.session.add(UserDepartmentAssignment(user_id=viewer.id, dept_id=employee.dept_id))
    db.session.commit()
    with backup_app.test_request_context('/?month=2026-06'):
        g.current_user = viewer
        assert [row['emp_no'] for row in employees_list().get_json()] == ['E1']


@pytest.mark.parametrize('kind', ['employee', 'department', 'shift'])
def test_create_current_reference_reuses_inactive_business_key(backup_app, kind):
    from routes.admin_core import create_employee, create_department, create_shift
    if kind == 'shift':
        row = Shift(shift_no='S1', shift_name='旧班', time_slots=[], is_active=False)
        db.session.add(row)
    else:
        row = Employee.query.one() if kind == 'employee' else Department.query.one()
        row.is_active = False
    db.session.commit()
    old_id = row.id
    handlers = {'employee': (create_employee, {'emp_no': 'E1', 'name': '恢复人员'}),
                'department': (create_department, {'dept_no': 'D1', 'dept_name': '恢复部门'}),
                'shift': (create_shift, {'shift_no': 'S1', 'shift_name': '恢复班次', 'time_slots': []})}
    handler, payload = handlers[kind]
    with backup_app.test_request_context('/', method='POST', json=payload):
        response = handler()
        assert not isinstance(response, tuple)
        assert response.get_json()['status'] == 'ok'
    assert row.id == old_id and row.is_active is True


def test_monthly_correction_cli_is_explicit_and_reports_partial_quality(historical_months, backup_app, tmp_path):
    from services.monthly_reference_service import register_monthly_reference_commands, get_month_reference
    register_monthly_reference_commands(backup_app)
    changes = tmp_path / 'month-correction.json'
    changes.write_text('{"name": "修正甲"}')
    result = backup_app.test_cli_runner().invoke(args=['correct-month-reference', '--month', '2026-06',
        '--kind', 'employee', '--key', 'E1', '--payload-file', str(changes), '--category', 'attendance'])
    assert result.exit_code == 0, result.output
    assert 'partial' in result.output
    assert get_month_reference('2026-06', 'employee', 'E1')['name'] == '修正甲'
    assert get_month_reference('2026-07', 'employee', 'E1')['name'] == '甲'


def test_daily_correction_uses_historical_employee_type(historical_months, backup_app):
    from routes.admin_attendance_overrides import save_daily_attendance_override_batch_api
    from models.daily_attendance_override import DailyAttendanceOverride
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    with backup_app.test_request_context('/', method='POST', json={
            'emp_id': employee.id, 'month': '2026-06', 'dates': ['2026-06-04'], 'status': '事假'}):
        g.current_user = User.query.filter_by(username='admin').one()
        response = save_daily_attendance_override_batch_api()
        assert not isinstance(response, tuple)
        assert response.get_json()['calendar']['employee']['name'] == '甲'
    assert DailyAttendanceOverride.query.one().status == '事假'


def test_manager_source_setting_change_does_not_change_historical_recalculation(backup_app):
    from services.manager_attendance_service import build_manager_rows, ManagerAttendanceOptions
    from models.dingtalk_sync_run import DingTalkSyncRun
    from models.system_setting import SystemSetting
    employee = Employee.query.one()
    employee.is_manager = True
    db.session.add(MonthlyReport(emp_id=employee.id, report_month='2026-06', manager_raw_data={'出勤天数': 20}))
    db.session.add(SystemSetting(key='manager_attendance_source', value='local'))
    db.session.commit()
    ensure_month_reference('2026-06')
    db.session.commit()
    SystemSetting.query.one().value = 'dingtalk'
    db.session.add(DingTalkSyncRun(account_set_id=1, month='2026-06', source='dingtalk', status='success'))
    db.session.commit()
    assert build_manager_rows(ManagerAttendanceOptions(month='2026-06'))[0]['attendance_days'] == 20


def test_sync_captures_only_target_month_in_the_business_transaction(backup_app):
    from services.card_attendance_sync_service import sync_card_attendance
    from tests.test_card_attendance_sync_service import FakeCardClient
    from services.monthly_reference_service import get_month_reference
    result = sync_card_attendance(1, '2026-06', FakeCardClient([
        {'person_no': 'E1', 'person_name': '甲', 'dept_name': '生产', 'card_no': '',
         'record_date': date(2026, 6, 5), 'times': ['08:00', '17:00']}]))
    assert result['status'] == 'success'
    assert get_month_reference('2026-06', 'employee', 'E1')['name'] == '甲'
    assert get_month_reference('2026-07', 'employee', 'E1') is None


def test_daily_import_batches_reference_reads_per_month(backup_app):
    from services.import_service import ImportService
    from sqlalchemy import event
    ensure_month_reference('2026-06')
    db.session.commit()
    reference_reads = []
    def record_query(conn, cursor, statement, parameters, context, executemany):
        if statement.lstrip().upper().startswith('SELECT') and 'monthly_reference_snapshots' in statement:
            reference_reads.append(statement)
    event.listen(db.engine, 'before_cursor_execute', record_query)
    try:
        result = ImportService._import_daily_records([['人员编号', '考勤日期', '实出勤小时'],
            *[['E1', f'2026-06-{day:02d}', 8] for day in range(1, 31)]])
        assert result['imported'] == 30
        assert len(reference_reads) < 10
    finally:
        event.remove(db.engine, 'before_cursor_execute', record_query)


@pytest.mark.parametrize('kind', ['employee', 'department', 'shift'])
def test_reference_number_change_cannot_disconnect_captured_history(historical_months, backup_app, kind):
    from routes.admin_core import update_employee, update_department, update_shift
    employee, _, shift = historical_months
    department = Department.query.filter_by(dept_no='D1').one()
    handlers = {'employee': (update_employee, employee, {'emp_no': 'NEW', 'name': '甲'}),
                'department': (update_department, department, {'dept_no': 'NEW', 'dept_name': '生产'}),
                'shift': (update_shift, shift, {'shift_no': 'NEW', 'shift_name': '白班'})}
    handler, row, payload = handlers[kind]
    with backup_app.test_request_context('/', method='PUT', json=payload):
        response = handler(row.id)
        assert isinstance(response, tuple) and response[1] == 409
    assert getattr(row, {'employee': 'emp_no', 'department': 'dept_no', 'shift': 'shift_no'}[kind]) != 'NEW'


def test_employee_delete_cannot_cascade_captured_business_history(historical_months, backup_app):
    from routes.admin_core import delete_employee
    employee, _, _ = historical_months
    with backup_app.test_request_context('/', method='DELETE'):
        response = delete_employee(employee.id)
        assert isinstance(response, tuple) and response[1] == 409
    assert DailyRecord.query.filter_by(emp_id=employee.id).count() == 2


def test_unbound_department_cleanup_preserves_snapshot_binding(historical_months, backup_app):
    from routes.admin_core import delete_unbound_departments
    employee, _, _ = historical_months
    warehouse = Department.query.filter_by(dept_no='D2').one()
    employee.dept_id = None
    db.session.commit()
    with backup_app.test_request_context('/', method='DELETE'):
        response = delete_unbound_departments()
        assert response.get_json()['skipped_history_bound'] >= 1
    assert db.session.get(Department, warehouse.id) is not None


def test_archive_inspection_is_read_only_and_lists_sourced_conflicting_candidates(backup_app, tmp_path):
    from models.account_set import AccountSetImport
    from models.monthly_reference_snapshot import MonthlyReferenceSnapshot
    from services.monthly_reference_service import scan_legacy_month_references
    archive = tmp_path / '2026_6月员工日报.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active.append(['人员编号', '人员名称', '考勤日期', '部门名称'])
    workbook.active.append(['E1', '档案甲', '2026-06-03', '生产'])
    workbook.save(archive)
    DailyRecord.query.one().raw_data = {'姓名': '日报甲'}
    db.session.add(AccountSetImport(account_set_id=1, source_filename=archive.name,
        stored_path=str(archive), file_type='daily', status='ok'))
    db.session.commit()
    report = scan_legacy_month_references(inspect_archives=True)['months'][0]
    names = next(row for row in report['evidence'] if row['field'] == 'name')
    assert names['conflict'] is True
    assert {row['value'] for row in names['candidates']} == {'日报甲', '档案甲'}
    assert report['archives'][0]['status'] == 'inspected'
    assert MonthlyReferenceSnapshot.query.count() == 0


def test_manager_effective_source_is_independent_for_each_month_member(backup_app):
    from services.manager_attendance_service import build_manager_rows, ManagerAttendanceOptions
    from models.dingtalk_sync_run import DingTalkSyncRun
    from models.monthly_reference_snapshot import MonthlyReferenceSnapshot
    first = Employee.query.one()
    first.is_manager = True
    second = Employee(emp_no='M2', name='经理二', is_manager=True)
    db.session.add(second)
    db.session.flush()
    for employee in (first, second):
        db.session.add(MonthlyReport(emp_id=employee.id, report_month='2026-06', manager_raw_data={'出勤天数': 20}))
    db.session.add(DingTalkSyncRun(account_set_id=1, month='2026-06', source='dingtalk', status='success'))
    db.session.commit()
    ensure_month_reference('2026-06')
    snapshot = MonthlyReferenceSnapshot.query.filter_by(month='2026-06', kind='employee', business_key='E1').one()
    snapshot.payload = {**snapshot.payload, 'manager_attendance_source': 'dingtalk'}
    db.session.commit()
    rows = build_manager_rows(ManagerAttendanceOptions(month='2026-06'))
    assert next(row for row in rows if row['emp_no'] == 'M2')['attendance_days'] == 20


def test_account_recalculation_uses_frozen_ingestion_source(backup_app, monkeypatch, tmp_path):
    from routes.admin_core import calculate_account_set
    from models.system_setting import SystemSetting
    from tests.test_dingtalk_manager_attendance_service import FakeDingTalkClient
    ensure_month_reference('2026-06')
    db.session.add(SystemSetting(key='manager_attendance_source', value='dingtalk'))
    db.session.commit()
    backup_app.config['CALC_PROGRESS_DIR'] = str(tmp_path / 'progress')
    # Replace only the external client; run the real recalculation route/service.
    monkeypatch.setattr('services.dingtalk_client.DingTalkClient', lambda: FakeDingTalkClient())
    with backup_app.test_request_context('/?mode=manager'):
        result = calculate_account_set(1)
        response, status = result if isinstance(result, tuple) else (result, result.status_code)
        assert status == 400
        assert response.get_json()['message'] == '该账套暂无可计算文件'


def test_query_bootstrap_uses_selected_month_membership(historical_months, backup_app):
    from routes.api_query import bootstrap
    employee, manager, shift = historical_months
    change_current(employee, manager, shift)
    with backup_app.test_request_context('/?month=2026-07'):
        g.current_user = User.query.filter_by(username='admin').one()
        result = bootstrap.__wrapped__().get_json()
        row = next(row for row in result['employees'] if row['emp_no'] == 'E1')
        assert (row['name'], row['dept_name'], row['is_manager']) == ('甲', '仓储', False)


def test_current_admin_bootstrap_excludes_inactive_candidates(backup_app):
    from routes.api_admin import bootstrap
    Department.query.one().is_active = False
    db.session.add(Shift(shift_no='S1', shift_name='退出班次', time_slots=[], is_active=False))
    db.session.commit()
    with backup_app.test_request_context('/'):
        g.current_user = User.query.filter_by(username='admin').one()
        result = bootstrap.__wrapped__().get_json()
        assert result['departments'] == [] and result['shifts'] == []
