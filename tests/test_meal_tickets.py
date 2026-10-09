import io
import tempfile
import unittest
from datetime import date, timedelta

from flask import Flask
from openpyxl import load_workbook, Workbook
import xlrd

from models import db
from models.account_set import AccountSet
from models.employee import Employee
from models.employee_attendance_override import EmployeeAttendanceOverride
from models.user import User, UserEmployeeAssignment
from routes import register_routes
from routes.auth_helpers import generate_token
from tests.csrf_helper import attach_origin


class MealTicketTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SECRET_KEY='meal-tests-secret',
            SQLALCHEMY_DATABASE_URI='sqlite:///' + self.tmp.name + '/test.db',
            SQLALCHEMY_TRACK_MODIFICATIONS=False, JWT_EXPIRES_DELTA=timedelta(hours=12),
            FRONTEND_ORIGIN='http://localhost:5173', UPLOAD_FOLDER=self.tmp.name,
            BACKUP_PRIVATE_DIR=self.tmp.name + '/private')
        db.init_app(self.app)
        register_routes(self.app)
        with self.app.app_context():
            db.create_all()
            admin = User(username='admin', role='admin', password_hash='unused')
            viewer = User(username='viewer', role='readonly', password_hash='unused',
                          page_permissions={'meal_ticket_query': True})
            emp = Employee(emp_no='001', name='员工甲')
            other = Employee(emp_no='002', name='员工乙')
            account = AccountSet(month='2026-08', name='八月', is_locked=True)
            db.session.add_all([admin, viewer, emp, other, account])
            db.session.flush()
            db.session.add_all([
                EmployeeAttendanceOverride(emp_id=emp.id, month='2026-08', actual_attendance_days=22),
                EmployeeAttendanceOverride(emp_id=other.id, month='2026-08', actual_attendance_days=0),
                UserEmployeeAssignment(user_id=viewer.id, emp_id=emp.id)])
            db.session.commit()
            self.emp_id, self.other_id, self.admin_id = emp.id, other.id, admin.id
            self.admin_token, self.viewer_token = generate_token(admin), generate_token(viewer)
        self.client = attach_origin(self.app.test_client())
        self.headers = {'Authorization': 'Bearer ' + self.admin_token}

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()
        self.tmp.cleanup()

    def post(self, path, data):
        return self.client.post('/api/meal-tickets' + path, json=data, headers=self.headers)

    def generate(self):
        r = self.post('/generate', {'recharge_month': '2026-09'})
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def confirm(self, batch):
        r = self.post('/confirm', {'batch_id': batch['id'], 'version': batch['version']})
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def adjustment(self, batch, amount):
        r = self.post('/adjustments', {'batch_id': batch['id'], 'version': batch['version'],
            'item_id': batch['items'][0]['id'], 'amount': amount, 'reason': '额外情况'})
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def test_navigation_places_meal_tickets_immediately_after_query_center(self):
        response = self.client.get('/api/query/navigation', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        slugs = [module['slug'] for module in response.get_json()['modules']]
        self.assertEqual(slugs[slugs.index('query') + 1], 'meal-tickets')

    def test_recharge_export_requires_confirmation(self):
        url = '/api/meal-tickets/export-recharge?recharge_month=2026-09'
        self.assertEqual(self.client.get(url, headers=self.headers).status_code, 404)
        self.generate()
        self.assertEqual(self.client.get(url, headers=self.headers).status_code, 409)

    def test_recharge_export_is_two_column_xls_with_remaining_amount(self):
        batch = self.confirm(self.adjustment(self.generate(), '16.50'))
        response = self.post('/payments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'176', 'kind':'recharge',
            'date':'2026-09-05', 'reference':'充值凭证', 'request_key':'export-partial'})
        self.assertEqual(response.status_code, 200)
        response = self.client.get('/api/meal-tickets/export-recharge?recharge_month=2026-09', headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/vnd.ms-excel')
        self.assertIn('.xls', response.headers['Content-Disposition'])
        self.assertTrue(response.data.startswith(bytes.fromhex('d0cf11e0a1b11ae1')))
        book = xlrd.open_workbook(file_contents=response.data)
        self.assertEqual(book.nsheets, 1)
        sheet = book.sheet_by_index(0)
        self.assertEqual((sheet.nrows, sheet.ncols), (2, 2))
        self.assertEqual(sheet.row_values(0), ['员工编号', '充值金额'])
        self.assertEqual(sheet.row_values(1), ['001', 16.5])
        self.assertEqual(sheet.cell_type(1, 0), xlrd.XL_CELL_TEXT)
        self.assertEqual(sheet.cell_type(1, 1), xlrd.XL_CELL_NUMBER)
        current = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(len(current['items'][0]['payments']), 1)

    def test_recharge_export_omits_settled_and_refund_items(self):
        batch = self.confirm(self.generate())
        response = self.post('/payments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'176', 'kind':'recharge',
            'date':'2026-09-05', 'reference':'充值凭证', 'request_key':'export-settled'})
        self.assertEqual(response.status_code, 200)
        url = '/api/meal-tickets/export-recharge?recharge_month=2026-09'
        for amount in (None, '-8'):
            if amount:
                self.adjustment(response.get_json(), amount)
            exported = self.client.get(url, headers=self.headers)
            self.assertEqual(exported.status_code, 200)
            self.assertEqual(xlrd.open_workbook(file_contents=exported.data).sheet_by_index(0).nrows, 1)

    def test_recharge_export_respects_personnel_permissions(self):
        batch = self.confirm(self.generate())
        adjusted = self.post('/adjustments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][1]['id'], 'amount':'8', 'reason':'补发'})
        self.assertEqual(adjusted.status_code, 200)
        url = '/api/meal-tickets/export-recharge?recharge_month=2026-09'
        viewer = {'Authorization':'Bearer ' + self.viewer_token}
        exported = self.client.get(url, headers=viewer)
        self.assertEqual(exported.status_code, 200)
        sheet = xlrd.open_workbook(file_contents=exported.data).sheet_by_index(0)
        self.assertEqual(sheet.col_values(0), ['员工编号', '001'])
        self.assertEqual(self.client.get(url).status_code, 401)

    def test_final_field_next_month_and_no_rounding(self):
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 22.25
            db.session.commit()
        batch = self.generate()
        self.assertEqual(batch['month'], '2026-08')
        self.assertEqual(batch['recharge_month'], '2026-09')
        self.assertEqual(batch['items'][0]['days'], 22.25)
        self.assertEqual(batch['items'][0]['base_amount'], 178)
        self.assertEqual(batch['departments'][0]['due_amount'], 178)
        self.assertEqual(batch['rate_cents'], 800)
        self.assertEqual(batch['rule_version'], 'actual-days-v1')

    def test_adjustment_partial_payment_refund_and_idempotency(self):
        batch = self.adjustment(self.generate(), '16')
        batch = self.confirm(batch)
        body = {'batch_id': batch['id'], 'version': batch['version'], 'item_id': batch['items'][0]['id'],
                'amount': '176', 'kind': 'recharge', 'date': '2026-09-05',
                'reference': '凭证一', 'request_key': 'unique-payment-1'}
        paid = self.post('/payments', body)
        self.assertEqual(paid.status_code, 200, paid.get_json())
        self.assertEqual(paid.get_json()['items'][0]['difference'], 16)
        retry = self.post('/payments', body)
        self.assertEqual(retry.status_code, 200)
        self.assertEqual(len(retry.get_json()['items'][0]['payments']), 1)
        body['amount'] = '175'
        self.assertEqual(self.post('/payments', body).status_code, 409)
        batch = self.adjustment(paid.get_json(), '-24')
        self.assertEqual(batch['items'][0]['difference'], -8)
        body.update(version=batch['version'], amount='8', kind='refund', request_key='refund-1')
        batch = self.post('/payments', body).get_json()
        self.assertEqual(batch['items'][0]['difference'], 0)
        self.assertEqual(batch['items'][0]['paid_amount'], 168)

    def test_confirm_detects_changed_source_and_frozen_batch(self):
        batch = self.generate()
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 23
            db.session.commit()
        self.assertEqual(self.post('/confirm', {'batch_id': batch['id'], 'version': batch['version']}).status_code, 409)
        batch = self.confirm(self.generate())
        self.assertEqual(self.post('/generate', {'recharge_month': '2026-09'}).status_code, 409)
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 24
            db.session.commit()
        r = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers)
        self.assertTrue(r.get_json()['source_changed'])
        self.assertEqual(r.get_json()['items'][0]['base_amount'], 184)

    def test_unconfirm_restores_draft_without_losing_details(self):
        batch = self.adjustment(self.generate(), '16')
        batch = self.post('/participation', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][1]['id'], 'excluded':True, 'reason':'本月不发'}).get_json()
        batch = self.confirm(batch)
        response = self.post('/unconfirm', {'batch_id':batch['id'], 'version':batch['version']})
        self.assertEqual(response.status_code, 200, response.get_json())
        draft = response.get_json()
        self.assertEqual(draft['status'], 'draft')
        self.assertEqual(draft['version'], batch['version'] + 1)
        self.assertIsNone(draft['confirmed_by'])
        self.assertEqual(draft['items'], batch['items'])
        with self.app.app_context():
            from models.meal_ticket import MealTicketBatch
            stored = db.session.get(MealTicketBatch, batch['id'])
            self.assertIsNone(stored.confirmed_at)
            self.assertTrue(AccountSet.query.first().is_locked)
        self.assertEqual(self.client.get('/api/meal-tickets/export-recharge?recharge_month=2026-09',
            headers=self.headers).status_code, 409)
        response = self.post('/payments', {'batch_id':draft['id'], 'version':draft['version'],
            'item_id':draft['items'][0]['id'], 'amount':'8', 'kind':'recharge',
            'date':'2026-09-05', 'reference':'未重新确认', 'request_key':'unconfirmed-payment'})
        self.assertEqual(response.status_code, 400)
        response = self.post('/participation', {'batch_id':draft['id'], 'version':draft['version'],
            'item_id':draft['items'][1]['id'], 'excluded':False, 'reason':'恢复核算'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.get_json()['items'][1]['excluded'])

    def test_attendance_recalculation_preserves_payments_and_manual_adjustments(self):
        batch = self.confirm(self.adjustment(self.generate(), '16'))
        batch = self.post('/payments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'192', 'kind':'recharge',
            'date':'2026-09-05', 'reference':'已充值', 'request_key':'recalc-paid'}).get_json()
        payments = batch['items'][0]['payments']
        for days, change, due, difference in [(23, 8, 200, 8), (23, 0, 200, 8), (20, -24, 176, -16)]:
            with self.app.app_context():
                EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = days
                db.session.commit()
            body = {'batch_id':batch['id'], 'version':batch['version']}
            response = self.post('/attendance-recalculation/preview', body)
            self.assertEqual(response.status_code, 200, response.get_json())
            preview = response.get_json()
            self.assertEqual(preview['rows'][0]['amount'], change)
            self.assertEqual(preview['issues'], [])
            unchanged = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
            self.assertEqual(unchanged['items'][0]['payments'], payments)
            self.assertEqual(unchanged['version'], batch['version'])
            response = self.post('/attendance-recalculation', {**body, 'source_digest':preview['source_digest']})
            self.assertEqual(response.status_code, 200, response.get_json())
            batch = response.get_json()
            row = batch['items'][0]
            self.assertEqual((row['days'], row['base_amount']), (22, 176))
            self.assertEqual((row['due_amount'], row['difference']), (due, difference))
            self.assertEqual(row['payments'], payments)
            self.assertEqual(row['adjustments'][0]['amount'], 16)
            self.assertFalse(batch['source_changed'])
        self.assertEqual(len(batch['items'][0]['adjustments']), 3)

    def test_attendance_recalculation_rejects_stale_preview_and_drafts(self):
        batch = self.generate()
        body = {'batch_id':batch['id'], 'version':batch['version']}
        self.assertEqual(self.post('/attendance-recalculation/preview', body).status_code, 409)
        batch = self.confirm(batch)
        body['version'] = batch['version']
        preview = self.post('/attendance-recalculation/preview', body).get_json()
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 23
            db.session.commit()
        self.assertEqual(self.post('/attendance-recalculation', {**body,
            'source_digest':preview['source_digest']}).status_code, 409)
        unchanged = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(unchanged['items'][0]['adjustments'], [])

    def test_attendance_recalculation_preserves_exclusion_and_reports_new_people(self):
        batch = self.generate()
        batch = self.post('/participation', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'excluded':True, 'reason':'本月不发'}).get_json()
        batch = self.confirm(batch)
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 23
            db.session.commit()
        body = {'batch_id':batch['id'], 'version':batch['version']}
        preview = self.post('/attendance-recalculation/preview', body).get_json()
        self.assertEqual(preview['rows'][0]['amount'], 0)
        batch = self.post('/attendance-recalculation', {**body, 'source_digest':preview['source_digest']}).get_json()
        self.assertTrue(batch['items'][0]['excluded'])
        self.assertEqual(batch['items'][0]['due_amount'], 0)
        self.assertEqual(self.post('/unconfirm', {'batch_id':batch['id'],
            'version':batch['version']}).status_code, 409)
        with self.app.app_context():
            db.session.add(Employee(emp_no='003', name='新增员工'))
            db.session.commit()
        body['version'] = batch['version']
        preview = self.post('/attendance-recalculation/preview', body).get_json()
        self.assertTrue(preview['issues'])
        self.assertEqual(self.post('/attendance-recalculation', {**body,
            'source_digest':preview['source_digest']}).status_code, 409)
    def test_unconfirm_allows_recalculation_and_requires_fresh_confirmation(self):
        batch = self.confirm(self.adjustment(self.generate(), '16'))
        with self.app.app_context():
            AccountSet.query.first().is_locked = False
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).one().actual_attendance_days = 23
            db.session.commit()
        response = self.post('/unconfirm', {'batch_id':batch['id'], 'version':batch['version']})
        self.assertEqual(response.status_code, 200, response.get_json())
        draft = response.get_json()
        self.assertTrue(draft['source_changed'])
        self.assertEqual(draft['items'][0]['base_amount'], 176)
        rebuilt = self.generate()
        self.assertEqual(rebuilt['items'][0]['base_amount'], 184)
        self.assertEqual(rebuilt['items'][0]['adjustment_amount'], 16)
        self.assertEqual(self.post('/confirm', {'batch_id':rebuilt['id'], 'version':rebuilt['version']}).status_code, 400)
        with self.app.app_context():
            AccountSet.query.first().is_locked = True
            db.session.commit()
        confirmed = self.confirm(rebuilt)
        self.assertEqual(confirmed['items'][0]['due_amount'], 200)
        self.assertEqual(self.client.get('/api/meal-tickets/export-recharge?recharge_month=2026-09',
            headers=self.headers).status_code, 200)

    def test_unconfirm_rejects_draft_stale_version_and_readonly_user(self):
        draft = self.generate()
        self.assertEqual(self.post('/unconfirm', {'batch_id':draft['id'], 'version':draft['version']}).status_code, 409)
        batch = self.confirm(draft)
        self.assertEqual(self.post('/unconfirm', {'batch_id':batch['id'], 'version':draft['version']}).status_code, 409)
        body = {'batch_id':batch['id'], 'version':batch['version']}
        viewer = {'Authorization':'Bearer ' + self.viewer_token}
        self.assertEqual(self.client.post('/api/meal-tickets/unconfirm', json=body, headers=viewer).status_code, 403)
        self.assertEqual(self.client.post('/api/meal-tickets/unconfirm', json=body).status_code, 401)
        current = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual((current['status'], current['version']), ('confirmed', batch['version']))
        self.assertEqual(self.post('/unconfirm', body).status_code, 200)
        self.assertEqual(self.post('/unconfirm', body).status_code, 409)

    def test_unconfirm_rejects_any_payment_history_even_after_reversal(self):
        batch = self.confirm(self.generate())
        body = {'batch_id':batch['id'], 'version':batch['version'], 'item_id':batch['items'][0]['id'],
            'amount':'8', 'kind':'recharge', 'date':'2026-09-05', 'reference':'实际充值', 'request_key':'rollback-paid'}
        paid = self.post('/payments', body)
        self.assertEqual(paid.status_code, 200)
        batch = paid.get_json()
        self.assertEqual(self.post('/unconfirm', {'batch_id':batch['id'], 'version':batch['version']}).status_code, 409)
        body.update(version=batch['version'], kind='reversal', reversal_id=batch['items'][0]['payments'][0]['id'],
            request_key='rollback-reversed', reference='登记错误')
        reversed_payment = self.post('/payments', body)
        self.assertEqual(reversed_payment.status_code, 200)
        batch = reversed_payment.get_json()
        self.assertEqual(batch['items'][0]['paid_amount'], 0)
        self.assertEqual(self.post('/unconfirm', {'batch_id':batch['id'], 'version':batch['version']}).status_code, 409)
        current = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(current['status'], 'confirmed')
        self.assertEqual(current['version'], batch['version'])
        self.assertEqual(current['items'][0]['payments'], batch['items'][0]['payments'])

    def test_unconfirm_invalidates_inflight_payment_version(self):
        batch = self.confirm(self.generate())
        response = self.post('/unconfirm', {'batch_id':batch['id'], 'version':batch['version']})
        self.assertEqual(response.status_code, 200, response.get_json())
        response = self.post('/payments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'8', 'kind':'recharge',
            'date':'2026-09-05', 'reference':'旧页面', 'request_key':'rollback-stale-payment'})
        self.assertEqual(response.status_code, 409)
        current = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(current['status'], 'draft')
        self.assertEqual(current['items'][0]['payments'], [])

    def test_permissions_export_and_delete_guard(self):
        batch = self.generate()
        viewer = {'Authorization': 'Bearer ' + self.viewer_token}
        r = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=viewer)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(len(r.get_json()['items']), 1)
        self.assertEqual(self.client.post('/api/meal-tickets/generate', json={'recharge_month':'2026-09'}, headers=viewer).status_code, 403)
        exported = self.client.get('/api/meal-tickets/export?recharge_month=2026-09', headers=viewer)
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(book.sheetnames, ['人员明细', '部门汇总'])
        self.assertEqual(book['人员明细'].max_row, 4)
        self.assertEqual(self.client.delete(f'/api/admin/employees/{self.emp_id}', headers=self.headers).status_code, 409)

    def test_locked_source_required_and_missing_source_blocked(self):
        with self.app.app_context():
            db.session.add(Employee(emp_no='003', name='无来源'))
            db.session.commit()
        batch = self.generate()
        self.assertTrue(any(i['error'] for i in batch['items']))
        self.assertEqual(self.post('/confirm', {'batch_id':batch['id'], 'version':batch['version']}).status_code, 400)

    def test_employee_departure_verification_reads_live_record_without_changing_draft(self):
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).delete()
            db.session.commit()
        batch = self.generate()
        self.assertEqual(batch['items'][0]['employment_status'], 'active')
        self.assertIsNone(batch['items'][0]['resigned_at'])
        viewer = {'Authorization': 'Bearer ' + self.viewer_token}
        scoped = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=viewer).get_json()
        self.assertEqual(len(scoped['items']), 1)
        self.assertEqual(scoped['items'][0]['employment_status'], 'active')
        with self.app.app_context():
            db.session.get(Employee, self.emp_id).resigned_at = date(2026, 9, 1)
            db.session.commit()
        response = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers)
        result = response.get_json()
        person = result['items'][0]
        self.assertEqual(person['employment_status'], 'resigned')
        self.assertEqual(person['resigned_at'], '2026-09-01')
        self.assertEqual(result['version'], batch['version'])
        self.assertFalse(person['excluded'])
        self.assertEqual(person['error'], batch['items'][0]['error'])
        self.assertEqual(person['due_amount'], batch['items'][0]['due_amount'])
        self.assertEqual(person['participation_history'], [])
        scoped = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=viewer).get_json()
        self.assertEqual(scoped['items'], [])  # Existing query permissions exclude resigned employees.
        with self.app.app_context():
            db.session.get(Employee, self.emp_id).resigned_at = None
            db.session.commit()
        restored = self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json()
        self.assertEqual(restored['items'][0]['employment_status'], 'active')
        self.assertIsNone(restored['items'][0]['resigned_at'])

    def exclude(self, batch, item, excluded=True, reason='实际已离职，本月不发'):
        response = self.post('/participation', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':item['id'], 'excluded':excluded, 'reason':reason})
        self.assertEqual(response.status_code, 200, response.get_json())
        return response.get_json()

    def test_no_issue_bypasses_missing_source_and_preserves_history_on_recalculation(self):
        with self.app.app_context():
            EmployeeAttendanceOverride.query.filter_by(emp_id=self.emp_id).delete()
            db.session.commit()
        batch = self.generate()
        self.assertTrue(batch['items'][0]['error'])
        batch = self.exclude(batch, batch['items'][0])
        batch = self.generate()
        item = batch['items'][0]
        self.assertTrue(item['excluded'])
        self.assertEqual(item['due_amount'], 0)
        self.assertTrue(item['error'])
        self.assertEqual(item['participation_history'][0]['reason'], '实际已离职，本月不发')
        self.assertEqual(item['participation_history'][0]['operator'], 'admin')
        self.confirm(batch)

    def test_no_issue_zeros_base_and_adjustments_and_restore_returns_original_amount(self):
        batch = self.adjustment(self.generate(), '16')
        batch = self.exclude(batch, batch['items'][0])
        item = batch['items'][0]
        self.assertEqual((item['base_amount'], item['adjustment_amount'], item['due_amount']), (0, 0, 0))
        self.assertEqual(batch['departments'][0]['due_amount'], 0)
        self.assertEqual(item['original_base_amount'], 176)
        batch = self.exclude(batch, item, False, '核对后结算离职前出勤')
        self.assertEqual(batch['items'][0]['due_amount'], 192)
        self.assertEqual(len(batch['items'][0]['participation_history']), 2)

    def test_participation_requires_reason_admin_current_version_and_draft(self):
        batch = self.generate()
        body = {'batch_id':batch['id'], 'version':batch['version'], 'item_id':batch['items'][0]['id'],
                'excluded':True, 'reason':''}
        self.assertEqual(self.post('/participation', body).status_code, 400)
        body.update(reason='本月不发', excluded='true')
        self.assertEqual(self.post('/participation', body).status_code, 400)
        body['excluded'] = True
        viewer = {'Authorization':'Bearer ' + self.viewer_token}
        self.assertEqual(self.client.post('/api/meal-tickets/participation', json=body, headers=viewer).status_code, 403)
        batch = self.exclude(batch, batch['items'][0])
        self.assertEqual(self.post('/participation', body).status_code, 409)
        batch = self.confirm(batch)
        body.update(version=batch['version'], excluded=False)
        self.assertEqual(self.post('/participation', body).status_code, 409)

    def test_confirmed_no_issue_can_be_supplemented_without_reviving_old_adjustments(self):
        batch = self.adjustment(self.generate(), '16')
        batch = self.exclude(batch, batch['items'][0])
        self.assertEqual(self.post('/adjustments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'8', 'reason':'先恢复核算'}).status_code, 400)
        batch = self.adjustment(self.confirm(batch), '24')
        self.assertEqual(batch['items'][0]['due_amount'], 24)
        self.assertEqual(batch['items'][0]['adjustment_amount'], 24)

    def test_no_issue_survives_removed_employee_scope_and_export_and_backup(self):
        from datetime import date
        from services.account_set_backup_service import collect_backup
        batch = self.generate()
        batch = self.exclude(batch, batch['items'][0])
        with self.app.app_context():
            db.session.get(Employee, self.emp_id).resigned_at = date(2026, 7, 31)
            db.session.commit()
        batch = self.generate()
        self.assertTrue(batch['items'][0]['excluded'])
        self.assertIn('移出', batch['items'][0]['error'])
        self.confirm(batch)
        exported = self.client.get('/api/meal-tickets/export?recharge_month=2026-09', headers=self.headers)
        book = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(book['人员明细'].cell(4, 7).value, 0)
        self.assertIn('本月不发', book['人员明细'].cell(4, 10).value)
        with self.app.app_context():
            document = collect_backup(AccountSet.query.first().id)
            self.assertTrue(document['datasets']['meal_items'][0]['source']['participation_history'][0]['excluded'])

    def test_recalculation_can_resolve_existing_draft_when_all_people_leave_scope(self):
        from datetime import date
        batch = self.generate()
        for item in batch['items']:
            batch = self.exclude(batch, item)
        with self.app.app_context():
            for employee in Employee.query.all():
                employee.resigned_at = date(2026, 7, 31)
            db.session.commit()
        batch = self.generate()
        self.assertTrue(all(item['excluded'] for item in batch['items']))
        self.assertEqual(batch['departments'][0]['due_amount'], 0)
        self.confirm(batch)

    def test_historical_preview_confirm_and_department_not_added(self):
        book = Workbook()
        s = book.active
        s.title = '员工充值记录'
        s.append(['部门名称','人员编号','人员名称','实际充值金额'])
        s.append(['未分配部门','001','员工甲',176])
        t = book.create_sheet('菜票登记')
        t.append(['部门','金额','登记人','备注'])
        t.append(['未分配部门',176,None,None])
        stream = io.BytesIO()
        book.save(stream)
        r = self.client.post('/api/meal-tickets/imports', data={
            'file': (io.BytesIO(stream.getvalue()), 'history.xlsx'), 'month': '2026-08', 'month_kind': 'attendance'}, headers=self.headers)
        self.assertEqual(r.status_code, 200, r.get_json())
        preview = r.get_json()
        self.assertEqual(preview['person_total'], 176)
        self.assertEqual(preview['department_total'], 176)
        self.assertEqual(preview['status'], 'preview')
        compared = self.client.get(f"/api/meal-tickets/imports/{preview['id']}/comparison", headers=self.headers)
        self.assertEqual(compared.status_code, 200)
        self.assertEqual(compared.get_json()[0]['base_amount'], 176)
        self.assertEqual(compared.get_json()[0]['difference'], 0)
        self.assertIsNone(self.client.get('/api/meal-tickets?recharge_month=2026-09', headers=self.headers).get_json())
        result = self.post(f"/imports/{preview['id']}/confirm", {'rows': preview['rows']})
        self.assertEqual(result.status_code, 200, result.get_json())
        self.assertEqual(result.get_json()['person_total'], 176)
        batch = self.generate()
        self.assertEqual(batch['items'][0]['paid_amount'], 0)

    def test_backup_round_trip_preserves_meal_records_and_original_file(self):
        from services.account_set_backup_service import export_backup, read_backup
        batch = self.confirm(self.generate())
        with self.app.app_context():
            account = AccountSet.query.filter_by(month='2026-08').one()
            document = read_backup(export_backup(account.id))
            self.assertEqual(document['datasets']['meal_batches'][0]['status'], 'confirmed')
            self.assertEqual(len(document['datasets']['meal_items']), 2)
            self.assertEqual(document['datasets']['meal_items'][0]['base_cents'], 17600)

    def test_manager_punch_field_and_cross_year(self):
        from unittest.mock import patch
        with self.app.app_context():
            emp = db.session.get(Employee, self.emp_id)
            emp.is_manager = True
            db.session.commit()
        with patch('services.meal_ticket_service.build_manager_rows', return_value=[{'emp_id':self.emp_id, 'punch_days':7, 'actual_attendance_days':20}]), \
             patch('services.meal_ticket_service.selected_monthly_report_raw', return_value={'source':'present'}):
            batch = self.generate()
        self.assertEqual(batch['items'][0]['base_amount'], 56)
        from services.meal_ticket_service import shift_month
        self.assertEqual(shift_month('2027-01', -1), '2026-12')

    def test_reversal_is_once_and_preserves_original(self):
        batch = self.confirm(self.generate())
        body = {'batch_id':batch['id'], 'version':batch['version'], 'item_id':batch['items'][0]['id'],
                'amount':'176', 'kind':'recharge', 'date':'2026-09-05', 'reference':'原凭证', 'request_key':'p1'}
        batch = self.post('/payments',body).get_json()
        original = batch['items'][0]['payments'][0]['id']
        body.update(version=batch['version'], kind='reversal', reversal_id=original, amount='0', request_key='reverse1', reference='登记错误')
        batch = self.post('/payments',body).get_json()
        self.assertEqual(batch['items'][0]['paid_amount'], 0)
        self.assertTrue(batch['items'][0]['payments'][0]['reversed'])
        body.update(version=batch['version'], request_key='reverse2')
        self.assertEqual(self.post('/payments',body).status_code,409)

    def test_invalid_and_excess_payment_and_stale_version_rejected(self):
        batch = self.generate()
        old_version = batch['version']
        batch = self.adjustment(batch,'1.50')
        self.assertEqual(self.post('/adjustments', {'batch_id':batch['id'], 'version':old_version,
            'item_id':batch['items'][0]['id'], 'amount':'1', 'reason':'重复旧页面'}).status_code,409)
        self.assertEqual(self.post('/adjustments', {'batch_id':batch['id'], 'version':batch['version'],
            'item_id':batch['items'][0]['id'], 'amount':'NaN', 'reason':'无效'}).status_code,400)
        batch = self.confirm(batch)
        body = {'batch_id':batch['id'], 'version':batch['version'], 'item_id':batch['items'][0]['id'],
            'amount':'178', 'kind':'recharge', 'date':'2026-09-01', 'reference':'超发', 'request_key':'overpay'}
        self.assertEqual(self.post('/payments',body).status_code,400)

    def test_recalculation_preserves_adjustment_and_unlocked_blocks_confirmation(self):
        batch = self.adjustment(self.generate(),'16')
        batch = self.generate()
        self.assertEqual(batch['items'][0]['due_amount'],192)
        with self.app.app_context():
            AccountSet.query.first().is_locked = False
            db.session.commit()
        self.assertEqual(self.post('/confirm',{'batch_id':batch['id'], 'version':batch['version']}).status_code,400)

    def test_missing_formula_cache_can_be_corrected_but_requires_reason(self):
        book = Workbook(); s = book.active; s.title='员工充值记录'
        s.append(['部门名称','人员编号','人员名称','实际充值金额'])
        s.append(['未分配部门','001','员工甲','=22*8'])
        stream=io.BytesIO(); book.save(stream)
        r=self.client.post('/api/meal-tickets/imports',data={'file':(io.BytesIO(stream.getvalue()),'formula.xlsx'),
            'month':'2026-09','month_kind':'recharge'}, headers=self.headers)
        preview=r.get_json()
        self.assertEqual(preview['month'],'2026-08')
        self.assertIsNone(preview['rows'][0]['amount'])
        self.assertTrue(preview['rows'][0]['error'])
        preview['rows'][0]['amount']=176
        self.assertEqual(self.post(f"/imports/{preview['id']}/confirm",{'rows':preview['rows']}).status_code,400)
        preview['rows'][0]['correction_reason']='按原纸账核对金额'
        self.assertEqual(self.post(f"/imports/{preview['id']}/confirm",{'rows':preview['rows']}).status_code,200)

    def test_history_file_and_restore_round_trip(self):
        from services.account_set_backup_service import export_backup,read_backup
        from services.account_set_restore_service import build_preview,restore_backup
        from models.meal_ticket import MealTicketImport
        book=Workbook(); s=book.active; s.title='员工充值记录'
        s.append(['部门名称','人员编号','人员名称','实际充值金额'])
        s.append(['未分配部门','001','员工甲',176])
        stream=io.BytesIO(); book.save(stream)
        self.client.post('/api/meal-tickets/imports',data={'file':(io.BytesIO(stream.getvalue()),'original.xlsx'),
            'month':'2026-08','month_kind':'attendance'}, headers=self.headers)
        self.confirm(self.generate())
        with self.app.app_context():
            document=read_backup(export_backup(AccountSet.query.first().id))
            self.assertIn(stream.getvalue(),document['_files'].values())
            db.session.remove(); db.drop_all(); db.create_all()
            db.session.add(User(id=self.admin_id,username='admin',role='admin',password_hash='unused'))
            db.session.commit()
            options={'employees':True,'departments':True,'shifts':True}
            result=build_preview(document,options)
            choices={r['key']:'backup' for r in result['rows'] if r['enabled']}
            restore_backup(document,options,choices,result['fingerprint'],self.admin_id)
            self.assertEqual(MealTicketImport.query.one().status,'preview')
            from pathlib import Path
            self.assertEqual(Path(MealTicketImport.query.one().stored_path).read_bytes(),stream.getvalue())

    def test_restore_cannot_add_payment_to_different_retained_snapshot(self):
        from services.account_set_backup_service import collect_backup
        from services.account_set_restore_service import build_preview, validate_selection
        from models.meal_ticket import MealTicketBatch, MealTicketItem, MealTicketPayment
        batch=self.confirm(self.generate())
        self.post('/payments', {'batch_id':batch['id'],'version':batch['version'],'item_id':batch['items'][0]['id'],
            'amount':'176','kind':'recharge','date':'2026-09-01','reference':'原凭证','request_key':'restore-payment'})
        with self.app.app_context():
            document=collect_backup(AccountSet.query.first().id)
            MealTicketPayment.query.delete()
            MealTicketBatch.query.one().status='draft'
            MealTicketItem.query.filter_by(emp_id=self.emp_id).one().base_cents=16000
            AccountSet.query.first().is_locked=False
            db.session.commit()
            p=build_preview(document,{})
            self.assertTrue(validate_selection(document,{},p['rows'],{}))

    def test_unlock_and_equal_total_source_correction_are_visible(self):
        from models.daily_attendance_override import DailyAttendanceOverride
        from datetime import date
        batch=self.generate()
        with self.app.app_context():
            db.session.add(DailyAttendanceOverride(emp_id=self.emp_id,record_date=date(2026,8,1),remark='新的核对说明'))
            db.session.commit()
        self.assertEqual(self.post('/confirm',{'batch_id':batch['id'],'version':batch['version']}).status_code,409)
        self.confirm(self.generate())
        with self.app.app_context():
            AccountSet.query.first().is_locked=False; db.session.commit()
        self.assertTrue(self.client.get('/api/meal-tickets?recharge_month=2026-09',headers=self.headers).get_json()['source_changed'])

    def test_conflicting_year_requires_explicit_confirmation(self):
        book=Workbook(); s=book.active; s.title='员工充值记录'
        s.append(['2025年8月充值']); s.append(['部门名称','人员编号','人员名称','实际充值金额'])
        s.append(['未分配部门','001','员工甲',176]); stream=io.BytesIO(); book.save(stream)
        r=self.client.post('/api/meal-tickets/imports',data={'file':(io.BytesIO(stream.getvalue()),'2025年8月充值.xlsx'),
            'month':'2026-08','month_kind':'attendance'},headers=self.headers)
        preview=r.get_json()
        self.assertTrue(preview['rows'][0]['period_conflict'])
        self.assertEqual(self.post(f"/imports/{preview['id']}/confirm",{'rows':preview['rows']}).status_code,400)


if __name__ == '__main__':
    unittest.main()
